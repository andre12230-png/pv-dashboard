"""
Chargement / ecriture du CSV consolide "Releves-pv.csv".

Format du CSV (locale FR) :
    Date;Prod_Jour;Inj_Jour;Conso_reseau_Jour;Conso_HC;Conso_HP
    29/04/2022;10,7;;;;
    ...

- separateur ";", decimales ","
- dates DD/MM/YYYY
- granularite journaliere
- cases vides autorisees (donnee non disponible)

API publique :
    load_releves_pv(path)   -> (production_df, reseau_df) au format pipeline
    read_releves_raw(path)  -> liste de dicts pour l'editeur de saisie
    write_releves_raw(path, rows) -> reecrit le CSV en preservant le format
    normalise_date_saisie(texte)   -> valide une date tapee ('DD/MM/YYYY')
    normalise_nombre_saisie(texte) -> valide un kWh tape (format FR)
    parse_enedis_fichier(path) -> lit un releve du reseau : classeur .xlsx ou
                              .csv simple d'Enedis, ou .csv "suivi de
                              consommation" d'Octopus (avec detail HC/HP)
    parse_enphase_fichier(path) -> lit un rapport Enphase Enlighten (.csv ou .zip)
    charger_recharges_ve(path) -> kWh recharges A LA MAISON, par journee
                              (lecture du recharges.json de l'app recharges-ve)
    retirer_jour_en_cours(valeurs) -> enleve la journee en cours (incomplete)
    merge_import(rows, colonne, valeurs) -> fusion sans doublon par date
    resample_sum(df, freq)  -> helper pandas (resample par somme)
"""

from __future__ import annotations

import csv
import json
import os
import re
import shutil
import statistics
import sys
from datetime import date as _date
from typing import Iterable

import pandas as pd


def _first_matching(cols: Iterable[str], candidates: Iterable[str]) -> str | None:
    cols_lower = {c.lower().strip(): c for c in cols}
    for cand in candidates:
        if cand.lower() in cols_lower:
            return cols_lower[cand.lower()]
    return None


# ----------------------------------------------------------------------
# Helpers
# ----------------------------------------------------------------------

def resample_sum(df: pd.DataFrame, freq: str) -> pd.DataFrame:
    """Resample par somme. freq: 'D' (jour), 'MS' (mois), 'YS' (annee)."""
    if df.empty:
        return df
    return df.resample(freq).sum(min_count=1)


# ----------------------------------------------------------------------
# CSV consolide "Releves-pv.csv"
# ----------------------------------------------------------------------

_RELEVES_DATE_CANDIDATES = ("Date", "Jour", "Date_jour")
_RELEVES_PROD_CANDIDATES = ("Prod_Jour", "Production", "Prod")
_RELEVES_INJ_CANDIDATES = ("Inj_Jour", "Injection", "Inj")
_RELEVES_SOUT_CANDIDATES = (
    "Conso_reseau_Jour", "Conso_réseau_Jour", "Conso_Jour", "Soutirage", "Conso",
)
_RELEVES_HC_CANDIDATES = ("Conso_HC", "HC", "Heures_creuses")
_RELEVES_HP_CANDIDATES = ("Conso_HP", "HP", "Heures_pleines")


def _read_releves_csv(path: str) -> pd.DataFrame:
    """Lit le CSV consolide en gerant l'encodage et le separateur FR."""
    for encoding in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
        try:
            df = pd.read_csv(
                path, sep=";", decimal=",", encoding=encoding,
                dtype=str, keep_default_na=False,
            )
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError(f"Encodage non reconnu pour {path}")

    df.columns = [c.strip() for c in df.columns]
    return df


def load_releves_pv(path: str) -> tuple[pd.DataFrame, pd.DataFrame]:
    """
    Charge le CSV consolide et retourne (production, reseau).

    - production : DataFrame indexe par timestamp (jour) avec 'production_kwh'.
    - reseau     : DataFrame indexe par timestamp (jour) avec 'injection_kwh',
                   'soutirage_kwh', 'releve_incomplet' (1.0 les jours avec
                   production saisie mais sans releve d'injection), et
                   eventuellement 'soutirage_hp_kwh' / 'soutirage_hc_kwh'
                   si les colonnes HP/HC sont renseignees.
    """
    df = _read_releves_csv(path)

    date_col = _first_matching(df.columns, _RELEVES_DATE_CANDIDATES)
    prod_col = _first_matching(df.columns, _RELEVES_PROD_CANDIDATES)
    inj_col = _first_matching(df.columns, _RELEVES_INJ_CANDIDATES)
    sout_col = _first_matching(df.columns, _RELEVES_SOUT_CANDIDATES)
    hc_col = _first_matching(df.columns, _RELEVES_HC_CANDIDATES)
    hp_col = _first_matching(df.columns, _RELEVES_HP_CANDIDATES)

    if date_col is None or prod_col is None:
        raise ValueError(
            f"Releves-pv.csv non reconnu ({path}). "
            f"Colonnes trouvees : {list(df.columns)}"
        )

    def _to_num(series: pd.Series) -> pd.Series:
        s = series.astype(str).str.strip().str.replace(",", ".", regex=False)
        s = s.replace({"": None, "nan": None, "NaN": None})
        return pd.to_numeric(s, errors="coerce")

    out = pd.DataFrame()
    out["timestamp"] = pd.to_datetime(df[date_col], errors="coerce", dayfirst=True)
    out["production_kwh"] = _to_num(df[prod_col])
    out["injection_kwh"] = _to_num(df[inj_col]) if inj_col else pd.NA
    out["soutirage_jour_kwh"] = _to_num(df[sout_col]) if sout_col else pd.NA
    out["soutirage_hc_kwh"] = _to_num(df[hc_col]) if hc_col else pd.NA
    out["soutirage_hp_kwh"] = _to_num(df[hp_col]) if hp_col else pd.NA

    out = out.dropna(subset=["timestamp"]).set_index("timestamp").sort_index()
    out = out[~out.index.duplicated(keep="last")]

    # Soutirage total : preference au champ direct ; sinon HC+HP.
    sout_direct = out["soutirage_jour_kwh"]
    sout_hphc = out[["soutirage_hc_kwh", "soutirage_hp_kwh"]].sum(
        axis=1, min_count=1,
    )
    soutirage = sout_direct.where(sout_direct.notna(), sout_hphc)

    production = pd.DataFrame(
        {"production_kwh": out["production_kwh"].fillna(0.0)},
        index=out.index,
    )
    reseau = pd.DataFrame(
        {
            "injection_kwh": out["injection_kwh"].fillna(0.0),
            "soutirage_kwh": soutirage.fillna(0.0),
            # 1.0 quand le jour a une production saisie mais pas de releve
            # d'injection Enedis : l'autoconso vaut alors toute la production
            # (estimation haute), les vues le signalent.
            "releve_incomplet": (
                out["injection_kwh"].isna() & out["production_kwh"].notna()
            ).astype(float),
            # 1.0 quand aucun releve de consommation reseau n'existe pour ce
            # jour (ni total, ni HC/HP). Sans marqueur, le fillna ci-dessus
            # ferait passer ces jours pour une consommation reelle de 0 kWh :
            # facture sous-evaluee et autoproduction affichee a 100 %.
            "conso_absente": soutirage.isna().astype(float),
        },
        index=out.index,
    )
    # On expose le detail HP/HC uniquement quand au moins une valeur est connue,
    # pour que le calcul de cout reseau puisse l'utiliser quand il existe.
    if out["soutirage_hp_kwh"].notna().any() or out["soutirage_hc_kwh"].notna().any():
        reseau["soutirage_hp_kwh"] = out["soutirage_hp_kwh"]
        reseau["soutirage_hc_kwh"] = out["soutirage_hc_kwh"]

    return production, reseau


# Colonnes du CSV consolide, en preservant l'orthographe avec accent.
RELEVES_HEADER = [
    "Date", "Prod_Jour", "Inj_Jour",
    "Conso_réseau_Jour", "Conso_HC", "Conso_HP",
]


def creer_csv_vide(path: str) -> None:
    """Cree un fichier de releves ne portant que sa ligne d'en-tete.

    Sert au tout premier lancement : sans fichier, l'application refusait de
    demarrer et il fallait le fabriquer a la main, avec le bon separateur, le
    bon encodage et l'accent de "Conso_réseau_Jour". Trois pieges pour qui
    n'a jamais ouvert un CSV.

    N'ecrase jamais un fichier existant : les donnees passent avant tout.
    """
    if os.path.exists(path):
        return
    dossier = os.path.dirname(os.path.abspath(path))
    os.makedirs(dossier, exist_ok=True)
    with open(path, "w", encoding="utf-8", newline="") as fh:
        fh.write(";".join(RELEVES_HEADER) + "\n")


def _date_key(row: dict[str, str]) -> tuple[int, int, int]:
    """Cle de tri chronologique pour une ligne avec date FR DD/MM/YYYY."""
    try:
        d, m, y = row["Date"].split("/")
        return (int(y), int(m), int(d))
    except (ValueError, KeyError):
        return (0, 0, 0)


def read_releves_raw(path: str) -> list[dict[str, str]]:
    """
    Lit le CSV consolide en preservant les valeurs sous forme de chaines
    (utile pour un editeur : on ne perd pas la representation virgule/point).

    Retourne une liste de dicts {colonne -> valeur_str}, triee par date
    croissante. Les valeurs manquantes sont des chaines vides.

    Leve ValueError si une colonne du format n'a AUCUN alias correspondant
    dans l'en-tete : mieux vaut refuser de lire que rendre une colonne vide
    qu'un enregistrement effacerait ensuite pour de bon. Les variantes de
    nom documentees (accent ou non, "Production" pour "Prod_Jour"...)
    restent acceptees.
    """
    if not os.path.exists(path):
        return []

    df = _read_releves_csv(path)  # deja en str grace au dtype=str

    # Mapping nom officiel -> nom trouve (pour gerer les variantes d'accent)
    aliases = {
        "Date": _RELEVES_DATE_CANDIDATES,
        "Prod_Jour": _RELEVES_PROD_CANDIDATES,
        "Inj_Jour": _RELEVES_INJ_CANDIDATES,
        "Conso_réseau_Jour": _RELEVES_SOUT_CANDIDATES,
        "Conso_HC": _RELEVES_HC_CANDIDATES,
        "Conso_HP": _RELEVES_HP_CANDIDATES,
    }
    colmap: dict[str, str] = {}
    manquantes: list[str] = []
    for officiel, candidats in aliases.items():
        trouve = _first_matching(df.columns, candidats)
        if trouve is None:
            manquantes.append(officiel)
        else:
            colmap[officiel] = trouve

    # Une colonne dont AUCUN alias ne correspond etait simplement absente de
    # colmap, et la comprehension ci-dessous donnait alors "" a toutes les
    # lignes -- sans exception ni avertissement. L'application demarrait
    # normalement (load_releves_pv n'exige que Date et Prod_Jour), les vues
    # affichaient des chiffres plausibles, puis le premier Enregistrer
    # reecrivait le fichier a partir de ces chaines vides : des annees de
    # releves detruites, et l'anomalie d'en-tete effacee du meme coup.
    if manquantes:
        raise ValueError(
            f"{os.path.basename(path)} : colonne(s) introuvable(s) dans "
            f"l'en-tete : {', '.join(manquantes)}.\n"
            f"En-tetes lus dans le fichier : {list(df.columns)}.\n"
            "Aucun des noms acceptes ne correspond. Le fichier n'est pas lu, "
            "car une colonne lue vide serait effacee au premier "
            "enregistrement.\n"
            "Cause la plus frequente : le fichier a ete ouvert puis "
            "reenregistre par Excel, qui a abime les accents de l'en-tete."
        )

    rows: list[dict[str, str]] = []
    for _, r in df.iterrows():
        row = {
            officiel: str(r[colmap[officiel]]).strip() if officiel in colmap else ""
            for officiel in RELEVES_HEADER
        }
        if not row["Date"]:
            continue
        rows.append(row)

    # Tri chronologique par date FR (DD/MM/YYYY).
    rows.sort(key=_date_key)
    return rows


# ----------------------------------------------------------------------
# Import d'un export quotidien Enedis (consommation ou injection)
# ----------------------------------------------------------------------

_RE_DATE_ISO = re.compile(r"^(\d{4})-(\d{2})-(\d{2})")
_RE_DATE_FR = re.compile(r"^(\d{1,2})/(\d{1,2})/(\d{4})$")
_RE_NOMBRE = re.compile(r"^-?\d+(\.\d+)?$")


def _parse_date_any(s: str) -> str | None:
    """Date ISO (2026-06-01) ou FR (01/06/2026) -> 'DD/MM/YYYY', sinon None."""
    m = _RE_DATE_ISO.match(s)
    if m:
        y, mo, d = (int(x) for x in m.groups())
    else:
        m = _RE_DATE_FR.match(s)
        if not m:
            return None
        d, mo, y = (int(x) for x in m.groups())
    try:
        _date(y, mo, d)
    except ValueError:
        return None
    return f"{d:02d}/{mo:02d}/{y:04d}"


def dates_illisibles(path: str) -> list[str]:
    """
    Dates du CSV que pandas n'arrive pas a lire, dans l'ordre du fichier.

    load_releves_pv fait to_datetime(errors="coerce") puis dropna : une date
    illisible fait donc disparaitre la journee de TOUS les calculs, en
    silence -- alors qu'elle reste parfaitement visible dans l'onglet Saisie.
    L'ecart entre les deux est incomprehensible tant que personne ne le dit.

    Le CSV n'est pas modifie : c'est un constat, la correction se fait a la
    main dans la vue Saisie.
    """
    if not os.path.exists(path):
        return []
    df = _read_releves_csv(path)
    date_col = _first_matching(df.columns, _RELEVES_DATE_CANDIDATES)
    if date_col is None:
        return []

    brutes = df[date_col].astype(str).str.strip()
    # Meme lecture que load_releves_pv, pour compter exactement ce qu'elle
    # ecarte -- ni plus, ni moins.
    lues = pd.to_datetime(brutes, errors="coerce", dayfirst=True)
    return [
        brut for brut, lisible in zip(brutes, lues.notna(), strict=True)
        if brut and not lisible
    ]


def dernieres_saisies(path: str) -> dict[str, _date]:
    """
    Pour chaque colonne du CSV, la date du dernier jour reellement renseigne.

    Sert a dire d'un coup d'oeil ce qui est a jour et ce qui ne l'est pas.
    Les grandeurs n'avancent pas au meme rythme : la production vient
    d'Enphase (disponible le jour meme), l'injection et la consommation
    d'Enedis (deux jours de retard environ), le detail HC/HP d'Octopus
    plus tard encore. Une date unique pour tout le fichier ne dirait pas
    grand-chose.

    Une case vide ne compte pas : seule une valeur saisie fait foi.
    """
    dernieres: dict[str, _date] = {}
    for row in read_releves_raw(path):
        jour = _parse_date_any(row.get("Date", ""))
        if jour is None:
            continue
        d, m, y = (int(x) for x in jour.split("/"))
        quand = _date(y, m, d)
        for colonne in RELEVES_HEADER[1:]:
            if (row.get(colonne) or "").strip():
                if colonne not in dernieres or quand > dernieres[colonne]:
                    dernieres[colonne] = quand
    return dernieres


# ----------------------------------------------------------------------
# Bornes de plausibilite a la saisie
# ----------------------------------------------------------------------
# La faute de frappe la plus courante est la virgule oubliee : 251 tape au
# lieu de 25,1. Sans borne haute la valeur passe, et elle fausse ensuite
# toutes les moyennes sans que rien ne le signale.
#
# Production et injection : le plafond DEPEND de la taille de l'installation.
# Sur 6 kWc, le maximum observe en 4 ans de releves est 40,204 kWh le
# 01/07/2022 (et 35,408 kWh injectes le 02/07/2022), soit 6,7 kWh par kWc.
# Avec une marge de moitie : 10 kWh par kWc, ce qui redonne 60 kWh sur 6 kWc.
#
# Ce plafond etait fige a 60 : une installation de 9 kWc voyait ses journees
# d'ete refusees a la saisie, avec un message parlant de valeur improbable
# alors qu'elle etait vraie.
KWH_MAX_PAR_KWC = 10.0
PUISSANCE_PAR_DEFAUT_KWC = 6.0
MAX_KWH_JOUR = KWH_MAX_PAR_KWC * PUISSANCE_PAR_DEFAUT_KWC
# Consommation reseau : elle ne depend pas de l'installation mais du
# chauffage et de la recharge du vehicule, elle monte donc bien plus haut.
# Maximum observe : 80 kWh soutires le 22/01/2025, dont 50,22 kWh en heures
# creuses le 09/01/2026. Marge x1,9 environ -> 150 kWh.
#
# Un plafond unique a 60 refuserait 14 journees d'hiver REELLES du fichier :
# la vue Saisie deviendrait impossible a enregistrer des qu'elle en affiche
# une. D'ou deux plafonds, et non un seul.
MAX_KWH_JOUR_CONSO = 150.0

PLAFONDS_COLONNES = {
    "Prod_Jour": MAX_KWH_JOUR,
    "Inj_Jour": MAX_KWH_JOUR,
    "Conso_réseau_Jour": MAX_KWH_JOUR_CONSO,
    "Conso_HC": MAX_KWH_JOUR_CONSO,
    "Conso_HP": MAX_KWH_JOUR_CONSO,
}


def plafond_colonne(colonne: str, puissance_kwc: float | None = None) -> float:
    """Plafond de plausibilite d'une colonne officielle du CSV.

    `puissance_kwc` adapte le plafond de production et d'injection a la
    taille de l'installation (10 kWh par kWc). Sans elle, on garde celui
    d'une installation de 6 kWc.

    Colonne inconnue : on prend le plafond le plus permissif, mieux vaut
    laisser passer que refuser une valeur legitime.
    """
    plafond = PLAFONDS_COLONNES.get(colonne, MAX_KWH_JOUR_CONSO)
    if puissance_kwc and colonne in ("Prod_Jour", "Inj_Jour"):
        return max(float(puissance_kwc) * KWH_MAX_PAR_KWC, 1.0)
    return plafond


def normalise_date_saisie(texte: str, aujourdhui: _date | None = None) -> str:
    """
    Valide une date tapee dans la vue Saisie et la rend au format 'DD/MM/YYYY'.

    Strict volontairement : pandas acceptait '2026' (-> 01/01/2026) et
    '01/2026' (-> 01/01/2026), si bien qu'une frappe incomplete ecrasait
    silencieusement le releve du 1er janvier. On exige donc jour, mois et
    annee sur 4 chiffres (la forme ISO 2026-07-30 passe aussi).

    Leve ValueError avec un message lisible si la date n'est pas valide.
    """
    texte = (texte or "").strip()
    if not texte:
        raise ValueError("date manquante")
    jour = _parse_date_any(texte)
    if jour is None:
        raise ValueError(
            f"date '{texte}' incomplete ou mal formee "
            "(attendu JJ/MM/AAAA, par exemple 30/07/2026)")

    # Une date dans le futur n'est pas un releve : c'est une faute de frappe
    # sur l'annee. 2062 au lieu de 2026 faisait passer le fichier de 1572 a
    # 14721 lignes (resample("D") comble les jours manquants) et divisait la
    # production moyenne par neuf, sans le moindre avertissement.
    d, m, y = (int(x) for x in jour.split("/"))
    reference = aujourdhui or _date.today()
    if _date(y, m, d) > reference:
        raise ValueError(
            f"date '{jour}' dans le futur : nous sommes le "
            f"{reference:%d/%m/%Y}, un releve ne peut pas la porter")
    return jour


def normalise_nombre_saisie(texte: str, colonne: str = "",
                            plafond: float = MAX_KWH_JOUR) -> str:
    """
    Valide une valeur tapee dans la vue Saisie et la rend au format FR
    (virgule decimale), espaces retires. Une case vide reste vide : c'est
    autorise.

    La valeur rendue est celle qui sera ECRITE dans le CSV : elle doit donc
    etre relisible par _parse_nombre et par pandas.

    Refuse ce qui n'a pas de sens pour des kWh : texte, valeur negative,
    'inf' et 'nan' (que float() acceptait, ce qui contaminait ensuite tous
    les totaux avec des NaN), et valeur au-dessus du `plafond` de
    plausibilite de la colonne (voir plafond_colonne).

    Leve ValueError avec un message lisible.
    """
    texte = (texte or "").strip()
    if not texte:
        return ""
    ou = f" ({colonne})" if colonne else ""
    valeur = texte.replace(",", ".").replace(" ", "").replace(" ", "")
    try:
        nombre = float(valeur)
    except ValueError:
        raise ValueError(
            f"valeur '{texte}'{ou} : attendu un nombre, par exemple 24,8"
        ) from None
    if nombre != nombre or nombre in (float("inf"), float("-inf")):
        raise ValueError(f"valeur '{texte}'{ou} : nombre non exploitable")
    if nombre < 0:
        raise ValueError(
            f"valeur '{texte}'{ou} : un releve en kWh ne peut pas etre negatif")
    if nombre > plafond:
        raise ValueError(
            f"valeur '{texte}'{ou} : {_fmt_kwh_fr(nombre)} kWh depasse le "
            f"maximum plausible de {_fmt_kwh_fr(plafond)} kWh pour une "
            "journee (virgule oubliee ?)")
    # On rend la valeur NORMALISEE, pas le texte tape. Avant, '22, 1' etait
    # accepte puis ecrit tel quel dans le CSV : a la relecture, plus rien ne
    # savait le lire et la journee valait 0 kWh -- pire qu'une case vide, qui
    # elle est reperee comme absente et comblee.
    return _fmt_kwh_fr(nombre)


def _parse_nombre(s: str) -> float | None:
    s = s.strip().replace(" ", "").replace(" ", "").replace(",", ".")
    if not s or not _RE_NOMBRE.match(s):
        return None
    return float(s)


def _fmt_kwh_fr(v: float, decimales: int = 3) -> str:
    """12.571 -> '12,571' ; 9.0 -> '9' (format FR sans zeros inutiles)."""
    s = f"{v:.{decimales}f}".rstrip("0").rstrip(".")
    return s.replace(".", ",")


def parse_enedis_csv(path: str) -> tuple[str, dict[str, str]]:
    """
    Lit un export quotidien Enedis (mon-compte-particulier.enedis.fr).

    Le format varie selon les versions du site : lignes de metadonnees,
    dates ISO ou FR, valeurs en Wh ou kWh. Tout est detecte automatiquement :
    - lignes retenues : celles qui commencent par une date suivie d'un nombre ;
    - type de donnees (via le texte du fichier) :
        'injection' / 'production'   -> colonne Inj_Jour
        'consommation' / 'soutirage' -> colonne Conso_réseau_Jour
    - unite : '(en kWh)' / '(en Wh)' dans les entetes, sinon heuristique
      (une valeur quotidienne mediane > 200 est forcement en Wh).

    Retourne (colonne_cible, {date 'DD/MM/YYYY': valeur kWh en chaine FR}).
    Leve ValueError avec un message clair si le fichier n'est pas reconnu.
    """
    return parse_enedis_csv_texte(_lire_texte(path))


_MOTS_INJECTION = ("injection", "production")
_MOTS_CONSO = ("consommation", "soutirage")


def _entete_du_fichier(texte: str) -> str:
    """Les lignes qui precedent la premiere ligne de donnees 'date;valeur'.

    C'est la que Enedis annonce le type ('Injection quotidienne'). Chercher
    dans tout le fichier serait trompeur : le nom de l'export contient
    souvent les DEUX mots ('Consommation-Production'), et les libelles
    parasites d'une autre section suffiraient a faire basculer le verdict.
    """
    entete = []
    for ligne in texte.splitlines():
        cellules = [c.strip().strip('"') for c in re.split(r"[;\t]", ligne)]
        if len(cellules) >= 2 and _parse_date_any(cellules[0]) is not None:
            if any(_parse_nombre(c) is not None for c in cellules[1:]):
                break  # premiere ligne de donnees : l'entete s'arrete ici
        entete.append(ligne)
    return "\n".join(entete)


def _type_de_donnees(texte: str) -> str:
    """Decide si l'export porte sur l'injection ou sur la consommation.

    Regle : on lit l'en-tete en priorite. Si les deux familles de mots y
    figurent, on REFUSE de deviner plutot que de risquer d'ecrire une conso
    dans la colonne Injection (l'ancienne version faisait gagner
    'production' sans autre forme de proces).
    """
    for zone, ou in ((_entete_du_fichier(texte), "l'en-tete"),
                     (texte, "le fichier")):
        bas = zone.lower()
        inj = [m for m in _MOTS_INJECTION if m in bas]
        conso = [m for m in _MOTS_CONSO if m in bas]
        if inj and conso:
            raise ValueError(
                f"Type de donnees ambigu : {ou} mentionne a la fois "
                f"{' / '.join(inj)} et {' / '.join(conso)}. Impossible de "
                "savoir s'il s'agit de l'injection ou de la consommation. "
                "Utilisez le classeur Excel Enedis, qui indique clairement "
                "chaque grandeur, ou le CSV 'suivi de consommation' "
                "d'Octopus."
            )
        if inj:
            return "Inj_Jour"
        if conso:
            return "Conso_réseau_Jour"
    raise ValueError(
        "Type de donnees introuvable : l'en-tete du fichier ne mentionne ni "
        "'injection'/'production' ni 'consommation', et l'application refuse "
        "de deviner dans quelle colonne ranger ces valeurs. Tout export "
        "quotidien 'date + valeur' convient -- Enedis, un fournisseur ou un "
        "onduleur -- pourvu que son en-tete dise de quelle grandeur il "
        "s'agit."
    )


def parse_enedis_csv_texte(texte: str) -> tuple[str, dict[str, str]]:
    """Comme parse_enedis_csv, mais a partir du contenu deja lu."""
    bas = texte.lower()
    colonne = _type_de_donnees(texte)

    # Unite explicite dans les entetes ; None -> heuristique plus bas.
    if "(kwh" in bas or "en kwh" in bas:
        unite_kwh: bool | None = True
    elif "(wh" in bas or "en wh" in bas:
        unite_kwh = False
    else:
        unite_kwh = None

    brutes: list[tuple[str, float]] = []
    for ligne in texte.splitlines():
        if ";" in ligne:
            cellules = ligne.split(";")
        elif "\t" in ligne:
            cellules = ligne.split("\t")
        else:
            continue
        cellules = [c.strip().strip('"') for c in cellules]
        if len(cellules) < 2:
            continue
        jour = _parse_date_any(cellules[0])
        if jour is None:
            continue  # ligne de metadonnees / entete
        valeur = next(
            (v for v in (_parse_nombre(c) for c in cellules[1:]) if v is not None),
            None,
        )
        if valeur is None:
            continue  # ex. ligne 'date debut;date fin' des metadonnees
        brutes.append((jour, valeur))

    if not brutes:
        raise ValueError(
            "Aucune ligne 'date;valeur' reconnue. Attendu : un export "
            "quotidien Enedis (une date et un nombre par ligne, separes "
            "par des points-virgules)."
        )

    if unite_kwh is None:
        unite_kwh = statistics.median(v for _, v in brutes) < 200

    # Doublons internes au fichier importe : la derniere valeur gagne.
    valeurs: dict[str, str] = {}
    for jour, v in brutes:
        valeurs[jour] = _fmt_kwh_fr(v if unite_kwh else v / 1000.0)
    return colonne, valeurs


# ----------------------------------------------------------------------
# Import du "suivi de consommation" Octopus (detail heures creuses / pleines)
# ----------------------------------------------------------------------

# En-tetes du fichier "suivi_conso_<PRM>_<adresse>_<date>.csv", telecharge
# depuis l'espace client OCTOPUS -- et non chez Enedis, malgre le numero de
# PRM dans son nom : c'est le fournisseur qui l'edite, lui seul connaissant
# les prix (le fichier porte des colonnes en euros, qu'on ignore ici).
# C'etait longtemps le seul releve qui donnait le detail HC/HP jour par jour
# (le classeur .xlsx d'Enedis ne contient que le total) ; d'autres fournisseurs
# nomment ces colonnes autrement, et un utilisateur qui recopie ses chiffres a
# la main ecrit simplement "HC" et "HP". On accepte donc plusieurs libelles
# par colonne.
#
# La comparaison reste une EGALITE, apres normalisation : jamais un "contient",
# sinon "consommation hc (euros)" passerait pour des kWh.


def _normalise_entete(nom: str) -> str:
    """'  CONSO_HC ' -> 'conso hc' (minuscules, soulignements = espaces)."""
    return " ".join(nom.replace("_", " ").lower().split())


_ENTETES_CONSO = (
    "consommation (kwh)", "consommation", "conso (kwh)", "conso",
    "consommation totale (kwh)", "consommation totale",
    "consommation reseau (kwh)", "consommation réseau (kwh)",
    "conso reseau jour", "conso réseau jour", "conso reseau", "conso réseau",
)
_ENTETES_HC = (
    "consommation hc (kwh)", "consommation hc", "conso hc (kwh)", "conso hc",
    "hc (kwh)", "hc", "heures creuses (en kwh)",
    "heures creuses (kwh)", "heures creuses",
    "consommation heures creuses (kwh)", "consommation heures creuses",
)
_ENTETES_HP = (
    "consommation hp (kwh)", "consommation hp", "conso hp (kwh)", "conso hp",
    "hp (kwh)", "hp", "heures pleines (en kwh)",
    "heures pleines (kwh)", "heures pleines",
    "consommation heures pleines (kwh)", "consommation heures pleines",
)

_COLONNES_SUIVI_CONSO = {
    **{e: "Conso_réseau_Jour" for e in _ENTETES_CONSO},
    **{e: "Conso_HC" for e in _ENTETES_HC},
    **{e: "Conso_HP" for e in _ENTETES_HP},
}


def parse_octopus_suivi_conso(texte: str) -> dict[str, dict[str, str]] | None:
    """
    Lit le "suivi de consommation" d'Octopus : une ligne par jour, avec des
    colonnes nommees dont le detail heures creuses / heures pleines.

    Les libelles acceptes pour chaque colonne sont listes dans _ENTETES_HC,
    _ENTETES_HP et _ENTETES_CONSO : le fichier d'Octopus, mais aussi des
    en-tetes simples ("HC", "HP", "Heures creuses"...) pour qui recopie ses
    chiffres a la main ou vient d'un autre fournisseur.

    Retourne {colonne_cible: {date 'DD/MM/YYYY': valeur kWh en chaine FR}},
    ou None si aucune colonne HC/HP n'est reconnue (le fichier releve alors
    du parseur generique parse_enedis_csv). Leve ValueError si une seule des
    deux colonnes est presente.
    """
    entetes: dict[int, str] | None = None
    resultat: dict[str, dict[str, str]] = {}

    for ligne in texte.splitlines():
        cellules = [c.strip().strip('"') for c in ligne.split(";")]
        if len(cellules) < 2:
            continue

        if entetes is None:
            trouves = {
                i: _COLONNES_SUIVI_CONSO[_normalise_entete(nom)]
                for i, nom in enumerate(cellules)
                if _normalise_entete(nom) in _COLONNES_SUIVI_CONSO
            }
            cibles = set(trouves.values())
            # Sans le detail HC/HP, ce n'est pas ce format : on laisse la
            # main au parseur generique.
            if {"Conso_HC", "Conso_HP"} <= cibles:
                entetes = trouves
            elif cibles & {"Conso_HC", "Conso_HP"}:
                # Une seule des deux : le dire, plutot que de laisser le
                # lecteur generique ranger ces kWh en consommation reseau.
                if "Conso_HC" in cibles:
                    presente, manquante = "creuses", "pleines"
                    colonne_manquante = "HP"
                else:
                    presente, manquante = "pleines", "creuses"
                    colonne_manquante = "HC"
                raise ValueError(
                    f"Colonne {colonne_manquante} introuvable : ce fichier "
                    f"donne les heures {presente}, mais pas les heures "
                    f"{manquante}. Pour importer le detail, il faut les deux "
                    "colonnes, nommees par exemple 'Consommation HC (kWh)' "
                    "et 'Consommation HP (kWh)', ou simplement 'HC' et 'HP'."
                )
            continue

        jour = _parse_date_any(cellules[0])
        if jour is None:
            continue  # ligne de metadonnees ou de total
        for i, cible in entetes.items():
            if i >= len(cellules):
                continue
            valeur = _parse_nombre(cellules[i])
            if valeur is not None:
                resultat.setdefault(cible, {})[jour] = _fmt_kwh_fr(valeur)

    return resultat or None


# ----------------------------------------------------------------------
# Index de compteur -> consommations journalieres
# ----------------------------------------------------------------------

def _en_date(jour: str) -> _date:
    """'05/03/2026' -> date(2026, 3, 5)."""
    d, m, y = (int(x) for x in jour.split("/"))
    return _date(y, m, d)


def consommation_depuis_index(index: dict[str, float]) -> dict[str, str]:
    """
    Convertit des index de compteur (des cumuls) en consommations par jour.

    `index` : {date 'DD/MM/YYYY': index en kWh}. Un index ne fait que monter :
    la consommation d'une journee est la difference entre deux releves qui se
    suivent. L'index releve un matin cloturant la journee precedente, la
    difference entre le 02/01 et le 01/01 est rangee au 01/01.

    Deux cas sont ecartes plutot que devines :
    - un releve manquant (les deux dates ne se suivent pas d'un jour) : la
      difference couvrirait plusieurs journees, et rien ne dit comment la
      repartir. On laisse le trou ;
    - un index qui recule (compteur remplace, retour a zero) : la difference
      serait negative.

    Retourne {date: valeur kWh en chaine FR}, prete pour merge_import.
    """
    resultat: dict[str, str] = {}
    jours = sorted(index, key=_en_date)
    # strict=False : la deuxieme liste a un element de moins, c'est voulu.
    for precedent, suivant in zip(jours, jours[1:], strict=False):
        if (_en_date(suivant) - _en_date(precedent)).days != 1:
            continue
        ecart = index[suivant] - index[precedent]
        if ecart < 0:
            continue
        resultat[precedent] = _fmt_kwh_fr(ecart)
    return resultat


# ----------------------------------------------------------------------
# Export d'index quotidiens d'Enedis (.xlsx)
# ----------------------------------------------------------------------

# Telecharge dans l'espace client Enedis ("Export_<PRM>_Index_<periode>.xlsx").
# C'est le seul relevé du GESTIONNAIRE DE RESEAU qui donne les heures creuses
# et pleines : il est donc disponible a tout foyer francais, quel que soit son
# fournisseur. Mais il ne donne pas des consommations : ce sont les index du
# compteur, des cumuls -- d'ou consommation_depuis_index ci-dessus.
#
# Forme du fichier (relevee sur un vrai export, 18/09/2026) : quelques lignes
# de metadonnees, puis une ligne d'en-tetes "Date du tele-releve | Index
# totalisateur (en kWh) | Heures Pleines (en kWh) | Heures Creuses (en kWh) |
# Non parametre... | Heures Pleines Saison Basse (en kWh) | ...". Les colonnes
# HP/HC utiles sont celles du CALENDRIER FOURNISSEUR (l'offre de l'abonne) ;
# celles du CALENDRIER DISTRIBUTEUR, en saison basse / haute, decoupent la
# meme energie autrement : les confondre doublerait la consommation. D'ou une
# comparaison par EGALITE des libelles, jamais un "contient".
_ENTETES_INDEX_TOTAL = (
    "index totalisateur (en kwh)", "index totalisateur (kwh)",
    "index totalisateur",
)


def _charger_openpyxl():
    """Importe openpyxl, avec un message lisible s'il manque."""
    try:
        import openpyxl
    except ImportError as exc:
        raise ValueError(
            "La lecture d'un fichier Excel demande le module 'openpyxl' "
            "(installez-le avec : py -m pip install openpyxl)."
        ) from exc
    return openpyxl


def _index_d_une_feuille(ws, production: bool = False
                         ) -> dict[str, dict[str, str]] | None:
    """
    Lit une feuille d'export d'index Enedis.

    `production` : cette feuille porte les index de production, c'est-a-dire
    l'energie INJECTEE au reseau (le compteur ne mesure que le surplus qui
    sort). Son en-tete est le meme que celui de la consommation : seul le nom
    de la feuille les distingue, d'ou ce drapeau -- sans lui, des kWh injectes
    seraient ranges en soutirage.

    Retourne {colonne_cible: {date: kWh}} converti en consommations, ou None
    si cette feuille n'est pas un export d'index (en-tetes absents).
    """
    total = "Inj_Jour" if production else "Conso_réseau_Jour"
    colonnes: dict[int, str] = {}
    col_date = -1
    index: dict[str, dict[str, float]] = {}

    for row in ws.iter_rows(values_only=True):
        cellules = list(row)

        if not colonnes:
            # Tant que l'en-tete n'est pas trouve, tout est metadonnees.
            trouves: dict[int, str] = {}
            date_ici = -1
            for i, valeur in enumerate(cellules):
                if not isinstance(valeur, str):
                    continue
                nom = _normalise_entete(valeur)
                if nom in _ENTETES_INDEX_TOTAL:
                    trouves[i] = total
                elif nom in _ENTETES_HP and not production:
                    trouves[i] = "Conso_HP"
                elif nom in _ENTETES_HC and not production:
                    trouves[i] = "Conso_HC"
                elif date_ici < 0 and nom.startswith("date"):
                    date_ici = i
            # L'index totalisateur est la signature du format : sans lui, ce
            # classeur est un autre export (conso/production), pas celui-ci.
            if date_ici >= 0 and total in trouves.values():
                colonnes, col_date = trouves, date_ici
            continue

        if col_date >= len(cellules):
            continue
        jour = _cellule_date(cellules[col_date])
        if jour is None:
            continue  # ligne vide, ou pied de tableau
        for i, cible in colonnes.items():
            if i >= len(cellules):
                continue
            valeur = _cellule_nombre(cellules[i])
            if valeur is not None:  # 'NA' et '-' : donnee indisponible
                index.setdefault(cible, {})[jour] = valeur

    if not colonnes:
        return None
    return {cible: consommation_depuis_index(valeurs)
            for cible, valeurs in index.items()}


def parse_enedis_index_xlsx(path: str) -> dict[str, dict[str, str]] | None:
    """
    Lit un export d'index quotidiens d'Enedis et le convertit en
    consommations journalieres. Chez un producteur, le classeur porte aussi
    une feuille d'index de production : elle donne l'injection.

    Retourne None si le classeur n'est pas de ce type : l'appelant passe
    alors au lecteur du classeur conso/production habituel.
    """
    openpyxl = _charger_openpyxl()
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    resultat: dict[str, dict[str, str]] = {}
    reconnu = False
    try:
        # Un producteur a deux feuilles : la consommation (avec le detail
        # HP/HC) et la production, c'est-a-dire l'injection. On prend les
        # deux, chacune dans sa colonne.
        for nom in wb.sheetnames:
            trouve = _index_d_une_feuille(wb[nom], "prod" in nom.lower())
            if trouve is None:
                continue  # feuille de garde, ou autre chose
            reconnu = True
            resultat.update(trouve)
    finally:
        wb.close()
    return resultat if reconnu else None


# Feuilles reconnues dans le classeur Excel Enedis. Pour un producteur,
# la "production" mesuree par Enedis est l'energie injectee au reseau.
_FEUILLES_ENEDIS = (
    ("production", "Inj_Jour"),
    ("injection", "Inj_Jour"),
    ("consommation", "Conso_réseau_Jour"),
    ("soutirage", "Conso_réseau_Jour"),
)


def _cellule_date(c) -> str | None:
    """Cellule Excel -> 'DD/MM/YYYY' si c'est une date, sinon None."""
    if isinstance(c, _date):  # datetime.datetime herite de datetime.date
        return f"{c.day:02d}/{c.month:02d}/{c.year:04d}"
    if isinstance(c, str):
        return _parse_date_any(c.strip())
    return None


def _cellule_nombre(c) -> float | None:
    if isinstance(c, (int, float)) and not isinstance(c, bool):
        return float(c)
    if isinstance(c, str):
        return _parse_nombre(c)
    return None


def parse_enedis_xlsx(path: str) -> dict[str, dict[str, str]]:
    """
    Lit un export Excel Enedis ('..._Export_energie_Consommation-Production_...').

    Le classeur contient une feuille par grandeur ('Export Consommation
    Quotidienne', 'Export Production Quotidienne') : des lignes de
    metadonnees puis un tableau 'Date | Valeur (en kWh)'. Les jours 'NA'
    (donnees indisponibles) sont ignores.

    Retourne {colonne_cible: {date 'DD/MM/YYYY': valeur kWh en chaine FR}}.
    """
    try:
        import openpyxl
    except ImportError as exc:
        raise ValueError(
            "La lecture d'un fichier Excel demande le module 'openpyxl' "
            "(installez-le avec : py -m pip install openpyxl)."
        ) from exc

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    resultat: dict[str, dict[str, str]] = {}
    try:
        for nom in wb.sheetnames:
            colonne = next(
                (col for cle, col in _FEUILLES_ENEDIS if cle in nom.lower()),
                None)
            if colonne is None:
                continue  # ex. "Page d'accueil"
            brutes: list[tuple[str, float]] = []
            textes: list[str] = []
            for row in wb[nom].iter_rows(values_only=True):
                cellules = [c for c in row if c is not None]
                textes.extend(c for c in cellules if isinstance(c, str))
                if len(cellules) < 2:
                    continue
                jour = _cellule_date(cellules[0])
                if jour is None:
                    continue  # metadonnees ('Date de debut : ', ...)
                valeur = next(
                    (v for v in (_cellule_nombre(c) for c in cellules[1:])
                     if v is not None), None)
                if valeur is not None:
                    brutes.append((jour, valeur))
            if not brutes:
                continue
            bas = " ".join(textes).lower()
            if "(en kwh" in bas or "(kwh" in bas:
                en_kwh = True
            elif "(en wh" in bas or "(wh" in bas:
                en_kwh = False
            else:
                en_kwh = statistics.median(v for _, v in brutes) < 200
            valeurs: dict[str, str] = {}
            for jour, v in brutes:  # doublon interne : la derniere gagne
                valeurs[jour] = _fmt_kwh_fr(v if en_kwh else v / 1000.0)
            resultat[colonne] = valeurs
    finally:
        wb.close()

    if not resultat:
        raise ValueError(
            "Aucune feuille 'Consommation' ou 'Production' avec des lignes "
            "'date + valeur' dans ce classeur. Attendu : le classeur Excel "
            "d'un gestionnaire de reseau (Enedis) ou un export equivalent."
        )
    return resultat


def parse_enedis_fichier(path: str) -> dict[str, dict[str, str]]:
    """
    Point d'entree de l'import d'un releve du reseau, trois formats acceptes :
    - .xlsx : le classeur officiel d'Enedis (feuilles conso + production),
      ou son export d'index quotidiens (converti en consommations) ;
    - .csv "suivi de consommation" d'Octopus : conso + detail HC/HP ;
    - .csv simple : une seule grandeur (conso ou injection).

    Le nom garde "enedis" pour ne pas casser les appels existants : les deux
    fournisseurs passent par ici.

    Retourne {colonne_cible: {date 'DD/MM/YYYY': valeur kWh en chaine FR}}.
    """
    if path.lower().endswith((".xlsx", ".xlsm")):
        index = parse_enedis_index_xlsx(path)
        if index is not None:
            return index
        return parse_enedis_xlsx(path)
    texte = _lire_texte(path)
    suivi = parse_octopus_suivi_conso(texte)
    if suivi is not None:
        return suivi
    colonne, valeurs = parse_enedis_csv_texte(texte)
    return {colonne: valeurs}


# ----------------------------------------------------------------------
# Import d'un rapport Enphase (production)
# ----------------------------------------------------------------------

# Date suivie eventuellement d'une heure : "2026-07-14 06:15:00 +0200",
# "14/07/2026 06:15". On ne garde que le jour, les valeurs sont totalisees.
_RE_DATE_HEURE_FR = re.compile(r"^(\d{1,2})/(\d{1,2})/(\d{4})")
# Nombre au format anglo-saxon uniquement (point decimal) : utilise quand
# le fichier est separe par des virgules, sinon "30,1" serait coupe en deux.
_RE_NOMBRE_POINT = re.compile(r"^-?\d+(\.\d+)?$")


def _parse_jour_any(s: str) -> str | None:
    """Comme _parse_date_any mais accepte une heure derriere la date."""
    jour = _parse_date_any(s)
    if jour is not None:
        return jour
    m = _RE_DATE_HEURE_FR.match(s.strip())
    if not m:
        return None
    d, mo, y = (int(x) for x in m.groups())
    try:
        _date(y, mo, d)
    except ValueError:
        return None
    return f"{d:02d}/{mo:02d}/{y:04d}"


def _decoupe_ligne(ligne: str) -> tuple[list[str], bool]:
    """Decoupe une ligne CSV. Retourne (cellules, separateur_virgule)."""
    if ";" in ligne:
        sep, virgule = ";", False
    elif "\t" in ligne:
        sep, virgule = "\t", False
    elif "," in ligne:
        sep, virgule = ",", True
    else:
        return [], False
    return [c.strip().strip('"') for c in ligne.split(sep)], virgule


def parse_enphase_texte(texte: str) -> dict[str, str]:
    """
    Lit le contenu d'un rapport Enphase Enlighten (menu Systeme > Rapports,
    "Energie mensuelle" recu par email).

    Le rapport donne l'energie produite toutes les 15 minutes : les valeurs
    d'un meme jour sont **totalisees** pour obtenir la production du jour.
    Un rapport deja quotidien (une ligne par jour) fonctionne aussi.

    Tolerant sur la forme : lignes d'entete ignorees, separateur ';', ',' ou
    tabulation, dates ISO ou FR avec ou sans heure, valeurs en Wh ou kWh
    (unite lue dans l'entete, sinon devinee sur le total journalier).

    Retourne {date 'DD/MM/YYYY': production kWh en chaine FR}, arrondie au
    dixieme comme les valeurs affichees par Enlighten et le reste du CSV.
    """
    bas = texte.lower()
    if "(kwh" in bas or "en kwh" in bas:
        unite_kwh: bool | None = True
    elif "(wh" in bas or "en wh" in bas:
        unite_kwh = False
    else:
        unite_kwh = None

    totaux: dict[str, float] = {}
    for ligne in texte.splitlines():
        cellules, virgule = _decoupe_ligne(ligne)
        if len(cellules) < 2:
            continue
        jour = _parse_jour_any(cellules[0])
        if jour is None:
            continue  # entete ou ligne de metadonnees
        valeur = None
        for cell in cellules[1:]:
            if virgule:
                # Fichier separe par des virgules : seul le point est decimal.
                brut = cell.replace(" ", "").replace(" ", "")
                v = float(brut) if _RE_NOMBRE_POINT.match(brut) else None
            else:
                v = _parse_nombre(cell)
            if v is not None:
                valeur = v
                break
        if valeur is None:
            continue
        totaux[jour] = totaux.get(jour, 0.0) + valeur

    if not totaux:
        raise ValueError(
            "Aucune ligne 'date;valeur' reconnue. Attendu : le rapport "
            "Enphase 'Energie mensuelle' (Enlighten > Menu > Systeme > "
            "Rapports), recu par email."
        )

    if unite_kwh is None:
        # Un total journalier de production depasse toujours 200 en Wh et
        # reste bien en dessous en kWh (installation domestique).
        unite_kwh = statistics.median(totaux.values()) < 200

    return {
        jour: _fmt_kwh_fr(v if unite_kwh else v / 1000.0, decimales=1)
        for jour, v in totaux.items()
    }


def parse_enphase_fichier(path: str) -> dict[str, dict[str, str]]:
    """
    Point d'entree de l'import Enphase : fichier CSV/texte du rapport, ou
    l'archive .zip telle qu'elle arrive dans l'email (le premier fichier
    texte de l'archive est lu).

    Retourne {'Prod_Jour': {date 'DD/MM/YYYY': valeur kWh en chaine FR}},
    meme forme que parse_enedis_fichier.
    """
    if path.lower().endswith(".zip"):
        import zipfile

        with zipfile.ZipFile(path) as zf:
            # Les .csv d'abord : une archive peut aussi contenir un
            # lisez-moi .txt qui ne nous interesse pas.
            noms = sorted(
                (n for n in zf.namelist()
                 if n.lower().endswith((".csv", ".txt"))),
                key=lambda n: not n.lower().endswith(".csv"),
            )
            if not noms:
                raise ValueError(
                    "Cette archive ne contient aucun fichier .csv. "
                    "Attendu : un releve de production, par exemple le "
                    "rapport Enphase recu par email.")
            brut = zf.read(noms[0])
        texte = _decode_texte(brut)
    else:
        with open(path, "rb") as fh:
            texte = _decode_texte(fh.read())

    return {"Prod_Jour": parse_enphase_texte(texte)}


def _decode_texte(brut: bytes) -> str:
    """Decode un contenu texte en essayant les encodages usuels."""
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return brut.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("Encodage du fichier non reconnu.")


def _lire_texte(path: str) -> str:
    """Lit un fichier texte quel que soit son encodage."""
    with open(path, "rb") as fh:
        return _decode_texte(fh.read())


# ----------------------------------------------------------------------
# Recharges du vehicule electrique (application "recharges-ve")
# ----------------------------------------------------------------------

def charger_recharges_ve(path: str) -> pd.DataFrame:
    """
    Lit le fichier `recharges.json` de l'application recharges-ve et rend,
    par journee civile, l'energie rechargee A LA MAISON et son cout.

    Seules les recharges `lieu == "maison"` comptent : elles seules passent
    par le compteur de la maison. Les recharges exterieures sont payees
    ailleurs et n'ont rien a voir avec le releve Enedis.

    Une charge de nuit va typiquement de 23:56 a 05:26 : elle est donc
    **repartie entre les deux journees** au prorata du temps passe dans
    chacune, pour rester comparable aux releves quotidiens.

    Le cout vient du fichier (`cout_eur`), calcule par l'autre application
    sur les vraies plages horaires : le repartir au prorata du cout reseau
    total le surestimerait beaucoup, ces charges tombant presque toutes en
    heures creuses.

    Colonnes : recharge_ve_kwh, recharge_ve_eur. Le fichier appartient a
    l'autre application : il est seulement lu, jamais modifie. Un DataFrame
    vide est rendu si le fichier est absent ou illisible (l'app doit demarrer
    meme sans, par exemple sur un autre PC).
    """
    from datetime import datetime, time, timedelta

    vide = pd.DataFrame(columns=["recharge_ve_kwh", "recharge_ve_eur"])
    try:
        with open(path, "r", encoding="utf-8") as fh:
            contenu = json.load(fh)
    except (OSError, ValueError) as exc:
        print(f"[warn] recharges VE illisibles ({path}) : {exc}", file=sys.stderr)
        return vide

    recharges = contenu.get("recharges", []) if isinstance(contenu, dict) else contenu
    kwh_jour: dict[_date, float] = {}
    eur_jour: dict[_date, float] = {}

    def _ajouter(jour: _date, kwh: float, eur: float) -> None:
        kwh_jour[jour] = kwh_jour.get(jour, 0.0) + kwh
        eur_jour[jour] = eur_jour.get(jour, 0.0) + eur

    for r in recharges:
        if r.get("lieu") != "maison":
            continue
        try:
            kwh = float(r["kwh"])
            debut = datetime.fromisoformat(r["date"])
            fin = datetime.fromisoformat(r["fin"])
        except (KeyError, TypeError, ValueError):
            continue  # ligne incomplete : on l'ignore plutot que de tout perdre
        try:
            eur = float(r.get("cout_eur") or 0.0)
        except (TypeError, ValueError):
            eur = 0.0
        duree = (fin - debut).total_seconds()
        if duree <= 0:
            # Duree nulle ou incoherente : tout sur le jour de debut.
            _ajouter(debut.date(), kwh, eur)
            continue
        jour = debut.date()
        while jour <= fin.date():
            minuit = datetime.combine(jour, time.min)
            deb_j = max(debut, minuit)
            fin_j = min(fin, minuit + timedelta(days=1))
            secondes = (fin_j - deb_j).total_seconds()
            if secondes > 0:
                part = secondes / duree
                _ajouter(jour, kwh * part, eur * part)
            jour += timedelta(days=1)

    if not kwh_jour:
        return vide
    out = pd.DataFrame({
        "recharge_ve_kwh": pd.Series(kwh_jour),
        "recharge_ve_eur": pd.Series(eur_jour),
    }).sort_index()
    out.index = pd.to_datetime(out.index)
    return out


def retirer_jour_en_cours(valeurs: dict[str, str],
                          aujourdhui: _date | None = None) -> str | None:
    """
    Retire du dictionnaire le jour en cours (journee forcement incomplete :
    le rapport s'arrete a l'heure de sa generation). Modifie le dictionnaire
    et retourne la date retiree, ou None s'il n'y en avait pas.
    """
    jour = aujourdhui or _date.today()
    cle = f"{jour.day:02d}/{jour.month:02d}/{jour.year:04d}"
    return cle if valeurs.pop(cle, None) is not None else None


def merge_import(rows: list[dict[str, str]], colonne: str,
                 valeurs: dict[str, str]) -> tuple[list[dict[str, str]], int, int, int]:
    """
    Fusionne des valeurs importees {date: kwh_str} dans les lignes du CSV,
    sans jamais creer de doublon (fusion par date).

    La valeur importee REMPLACE celle de la colonne pour les dates deja
    presentes ; les autres colonnes de la ligne ne sont pas touchees.

    Retourne (lignes, nb_nouveaux, nb_remplaces, nb_inchanges).
    """
    toutes = {r["Date"]: dict(r) for r in rows}
    nb_nouveaux = nb_remplaces = nb_inchanges = 0
    for date_fr, val in valeurs.items():
        if date_fr in toutes:
            if toutes[date_fr].get(colonne, "") == val:
                nb_inchanges += 1
            else:
                nb_remplaces += 1
            toutes[date_fr][colonne] = val
        else:
            ligne = {c: "" for c in RELEVES_HEADER}
            ligne["Date"] = date_fr
            ligne[colonne] = val
            toutes[date_fr] = ligne
            nb_nouveaux += 1
    return list(toutes.values()), nb_nouveaux, nb_remplaces, nb_inchanges


def controle_apres_import(rows: list[dict[str, str]],
                          colonne: str,
                          valeurs: dict[str, str]) -> str | None:
    """
    Verifie qu'un import ne produit pas de valeurs physiquement impossibles.

    Seul controle vraiment discriminant : **l'injection ne peut jamais
    depasser la production du meme jour** (on n'injecte que du surplus).
    Un export de consommation range par erreur dans la colonne Injection
    se trahit ainsi immediatement, la conso etant sans rapport avec la
    production du jour.

    Retourne un message d'avertissement, ou None si tout est coherent.
    """
    if colonne != "Inj_Jour" or not valeurs:
        return None
    par_date = {r["Date"]: r for r in rows}

    def _num(v: str) -> float | None:
        v = (v or "").strip().replace(",", ".")
        try:
            return float(v) if v else None
        except ValueError:
            return None

    suspects, exemples = 0, []
    compares = 0
    for jour, val in valeurs.items():
        inj = _num(val)
        prod = _num(par_date.get(jour, {}).get("Prod_Jour", ""))
        if inj is None or prod is None or prod <= 0:
            continue
        compares += 1
        if inj > prod * 1.05:  # 5 % de tolerance (arrondis, pas de mesure)
            suspects += 1
            if len(exemples) < 3:
                exemples.append(f"{jour} : {val} injecte pour {prod:g} produit")
    if compares == 0 or suspects <= compares * 0.10:
        return None
    return (
        f"Attention : sur {compares} jour(s) comparables, {suspects} auraient "
        f"une injection SUPERIEURE a la production.\n\n"
        + "\n".join("  - " + e for e in exemples)
        + "\n\nC'est impossible : on n'injecte que le surplus. Ce fichier "
        "contient probablement une autre grandeur (une consommation ?) "
        "que l'application a prise pour de l'injection."
    )


class ReductionRefusee(ValueError):
    """L'ecriture demandee ferait disparaitre des donnees presentes sur le disque.

    Levee par write_releves_raw quand le nombre de lignes, ou le nombre de
    valeurs d'une colonne, baisserait par rapport au fichier existant.

    La vue Saisie l'attrape, decrit la perte a l'utilisateur, et ne recommence
    avec autoriser_reduction=True que s'il confirme. L'attribut `pertes` porte
    la description lisible, prete a etre affichee.
    """

    def __init__(self, pertes: str):
        self.pertes = pertes
        super().__init__(
            f"Enregistrement refuse : {pertes}. "
            "Si cet effacement est voulu, il doit etre confirme."
        )


def _volumetrie(rows: list[dict[str, str]]) -> tuple[int, dict[str, int]]:
    """(nombre de lignes datees, nombre de valeurs non vides par colonne)."""
    par_colonne = {col: 0 for col in RELEVES_HEADER[1:]}
    lignes = 0
    for row in rows:
        if not (row.get("Date") or "").strip():
            continue
        lignes += 1
        for col in RELEVES_HEADER[1:]:
            if (row.get(col) or "").strip():
                par_colonne[col] += 1
    return lignes, par_colonne


def _volumetrie_sur_disque(path: str) -> tuple[int, dict[str, int]]:
    """Volumetrie du fichier existant, (0, {}) s'il n'y en a pas.

    Si le fichier existe mais n'est plus lisible comme CSV structure (en-tete
    abime, encodage casse), on ne peut plus compter par colonne ; le nombre de
    lignes reste mesurable et suffit deja a reperer une troncature.
    """
    if not os.path.exists(path):
        return 0, {}
    try:
        return _volumetrie(read_releves_raw(path))
    except (OSError, ValueError):
        pass
    try:
        with open(path, encoding="utf-8", errors="replace") as fh:
            lignes = [ligne for ligne in fh.read().splitlines() if ligne.strip()]
        return max(len(lignes) - 1, 0), {}   # -1 pour la ligne d'en-tete
    except OSError:
        return 0, {}


def _enumere(morceaux: list[str]) -> str:
    """['a', 'b', 'c'] -> 'a, b et c'."""
    if len(morceaux) == 1:
        return morceaux[0]
    return ", ".join(morceaux[:-1]) + " et " + morceaux[-1]


def pertes_a_l_ecriture(path: str,
                        rows: list[dict[str, str]]) -> str | None:
    """
    Decrit ce qu'ecrire `rows` dans `path` ferait disparaitre, ou None.

    Compare au fichier existant sur le disque le nombre de lignes ET le
    nombre de valeurs non vides de chaque colonne. Les deux comptent : vider
    une colonne sur toutes les lignes ne change pas le nombre de lignes.

    Sert a decrire la perte a l'utilisateur AVANT de la lui faire subir ;
    write_releves_raw s'en sert pour refuser par defaut.
    """
    lignes_avant, cols_avant = _volumetrie_sur_disque(path)
    lignes_apres, cols_apres = _volumetrie(rows)

    pertes: list[str] = []
    if lignes_apres < lignes_avant:
        pertes.append(f"{lignes_avant - lignes_apres} ligne(s)")
    for col, avant in cols_avant.items():
        manque = avant - cols_apres.get(col, 0)
        if manque > 0:
            pertes.append(f"{manque} valeur(s) de {col}")
    if not pertes:
        return None
    return _enumere(pertes) + " vont etre supprimees"


def write_releves_raw(path: str, rows: list[dict[str, str]],
                      autoriser_reduction: bool = False) -> None:
    """
    Reecrit le CSV consolide en preservant le format FR :
    - separateur ";"
    - dates DD/MM/YYYY
    - decimales avec virgule
    - encodage UTF-8 (sans BOM)
    - en-tete avec accent : 'Conso_réseau_Jour'
    Les valeurs sont passees telles quelles (deja en chaines avec virgule),
    a ceci pres qu'une valeur contenant un ';' est mise entre guillemets pour
    rester relisible.

    Ecriture atomique : le contenu est ecrit dans un fichier temporaire puis
    substitue d'un coup via os.replace. Si l'ecriture echoue en cours de
    route (verrou de synchro cloud, disque plein, coupure), le CSV existant reste
    intact. L'ancienne version est conservee dans <path>.bak.

    Le fichier est reconstruit INTEGRALEMENT a partir de `rows` : ce qui n'y
    figure pas disparait. Une lecture partielle deviendrait donc une
    suppression definitive. D'ou le controle de volumetrie : si le nombre de
    lignes ou le nombre de valeurs d'une colonne baisse, ReductionRefusee est
    levee. `autoriser_reduction=True` passe outre -- l'appelant doit alors
    avoir fait confirmer la perte, que pertes_a_l_ecriture sait decrire.
    """
    if not autoriser_reduction:
        pertes = pertes_a_l_ecriture(path, rows)
        if pertes is not None:
            raise ReductionRefusee(pertes)

    rows_sorted = [row for row in sorted(rows, key=_date_key)
                   if row.get("Date")]

    tmp_path = path + ".tmp"
    # Instantane de la version actuelle, pris AVANT d'ecrire quoi que ce soit.
    # C'est lui qui deviendra le .bak et l'archive horodatee -- mais seulement
    # si la publication reussit. Deux conditions a tenir ensemble :
    #   - les sauvegardes doivent contenir ce qui etait sur le disque AVANT
    #     cette ecriture, donc on ne peut pas les prendre apres os.replace
    #     (a ce moment-la `path` porte deja le nouveau contenu) ;
    #   - elles ne doivent pas exister si os.replace echoue, donc on ne peut
    #     pas les poser avant.
    # D'ou l'instantane intermediaire : capture tot, promu tard.
    prev_path = path + ".prev.tmp"
    instantane = None
    try:
        if os.path.exists(path):
            try:
                shutil.copy2(path, prev_path)
                instantane = prev_path
            except OSError as exc:
                # Best effort : ne pas pouvoir sauvegarder ne doit pas
                # empecher d'enregistrer les nouvelles donnees.
                print(f"[warn] instantane de {path} impossible : {exc}",
                      file=sys.stderr)

        with open(tmp_path, "w", encoding="utf-8", newline="") as fh:
            # csv.writer met d'office entre guillemets toute valeur contenant
            # le separateur. Avec ";".join(cells), un ';' dans une valeur
            # produisait 7 champs au lieu de 6 et rendait le fichier illisible
            # pour de bon : ParserError au demarrage suivant, sans retour
            # possible. lineterminator est impose, sans quoi csv.writer
            # ecrirait des CRLF et changerait les fins de ligne du fichier
            # entier (le CSV reel est en LF).
            graveur = csv.writer(fh, delimiter=";", lineterminator="\n")
            graveur.writerow(RELEVES_HEADER)
            for row in rows_sorted:
                graveur.writerow([row.get(col, "") for col in RELEVES_HEADER])
            # NTFS journalise le RENOMMAGE, pas les donnees : sans ces deux
            # lignes, une coupure de courant juste apres peut publier un CSV
            # de 0 octet, le renommage ayant ete enregistre mais pas son
            # contenu. flush() vide le tampon de Python, fsync() force le
            # disque a ecrire pour de bon.
            fh.flush()
            os.fsync(fh.fileno())

        # Publication atomique. Si elle echoue (CSV ouvert dans Excel, verrou
        # de partage Windows, disque plein), on ressort par le finally :
        # aucune sauvegarde n'aura ete consommee, aucun temporaire ne restera,
        # et le CSV existant sera intact. Quinze tentatives ratees purgeaient
        # auparavant tout l'historique fin, en le remplacant par quinze copies
        # identiques du fichier courant.
        os.replace(tmp_path, path)

        # A partir d'ici seulement, l'ancienne version est promue.
        if instantane is not None:
            archiver_sauvegarde(path, source=instantane)
            try:
                os.replace(instantane, path + ".bak")
                instantane = None   # devenu le .bak : plus rien a nettoyer
            except OSError as exc:
                print(f"[warn] sauvegarde {path}.bak impossible : {exc}",
                      file=sys.stderr)
    finally:
        for reste in (tmp_path, prev_path):
            if os.path.exists(reste):
                try:
                    os.remove(reste)
                except OSError:
                    pass  # un fichier verrouille sera repris au prochain passage


# Rotation des sauvegardes, sur le modele de l'application recharges-ve :
# des copies rapprochees pour rattraper une fausse manipulation immediate,
# et une par jour pour remonter plus loin sans garder des milliers de
# fichiers. Le .bak reste a cote du CSV, c'est le filet le plus visible.
MAX_SAUVEGARDES = 15          # copies horodatees les plus recentes
MAX_SAUVEGARDES_JOUR = 30     # une par journee


def dossier_sauvegardes(path: str) -> str:
    return os.path.join(os.path.dirname(os.path.abspath(path)), "backups")


def archiver_sauvegarde(path: str, source: str | None = None) -> str | None:
    """
    Copie le CSV actuel dans backups/ avant qu'il ne soit remplace, et fait
    le menage. Deux familles de fichiers :

    - `<nom>_<AAAAMMJJ_HHMMSS>.csv` : les MAX_SAUVEGARDES plus recentes ;
    - `<nom>_jour_<AAAAMMJJ>.csv`   : la premiere de chaque journee, gardee
      MAX_SAUVEGARDES_JOUR jours. Elle echappe a la rotation ci-dessus.

    Sans elle, deux enregistrements successifs suffisaient a perdre la seule
    version d'avant : le .bak etait ecrase par le second.

    `path` sert a NOMMER les copies et a situer le dossier backups ; `source`
    est le fichier dont le CONTENU est archive. Les deux different au moment
    ou write_releves_raw archive : `path` porte deja la nouvelle version,
    l'instantane de l'ancienne est ailleurs. Par defaut, source = path.

    Retourne le chemin de la copie horodatee, ou None en cas d'echec (la
    sauvegarde ne doit jamais empecher l'enregistrement).
    """
    from datetime import datetime

    try:
        dossier = dossier_sauvegardes(path)
        os.makedirs(dossier, exist_ok=True)
        base = os.path.splitext(os.path.basename(path))[0]
        maintenant = datetime.now()

        # Deux enregistrements dans la meme seconde donneraient le meme nom :
        # la seconde copie ecraserait la premiere au lieu de s'ajouter.
        horodatage = f"{maintenant:%Y%m%d_%H%M%S}"
        copie = os.path.join(dossier, f"{base}_{horodatage}.csv")
        suffixe = 1
        while os.path.exists(copie):
            copie = os.path.join(dossier, f"{base}_{horodatage}_{suffixe}.csv")
            suffixe += 1
        a_archiver = source or path
        shutil.copy2(a_archiver, copie)

        # Une sauvegarde par journee, posee seulement si elle n'existe pas.
        du_jour = os.path.join(dossier, f"{base}_jour_{maintenant:%Y%m%d}.csv")
        if not os.path.exists(du_jour):
            shutil.copy2(a_archiver, du_jour)

        _faire_le_menage(dossier, f"{base}_2", MAX_SAUVEGARDES)
        _faire_le_menage(dossier, f"{base}_jour_", MAX_SAUVEGARDES_JOUR)
        return copie
    except OSError as exc:
        print(f"[warn] archivage de {path} impossible : {exc}", file=sys.stderr)
        return None


def _faire_le_menage(dossier: str, prefixe: str, combien: int) -> None:
    """Ne garde que les `combien` fichiers les plus recents du prefixe.

    Tri sur la date de modification et non sur le nom : `copy2` preserve la
    date du contenu, alors qu'un tri alphabetique melangerait les deux
    familles de noms.
    """
    fichiers = [
        os.path.join(dossier, n) for n in os.listdir(dossier)
        if n.startswith(prefixe) and n.endswith(".csv")
    ]
    fichiers.sort(key=os.path.getmtime, reverse=True)
    for vieux in fichiers[combien:]:
        try:
            os.remove(vieux)
        except OSError:
            pass  # un fichier verrouille sera repris au prochain passage
