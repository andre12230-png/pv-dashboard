"""
Lecture des fichiers a importer : releves du gestionnaire de reseau (CSV et
classeurs Enedis, export d'index), suivi de consommation du fournisseur
(detail HC/HP), rapport d'onduleur et rapport Enphase (production).

Chaque lecteur rend {colonne: {date 'DD/MM/YYYY': valeur kWh en chaine FR}} ;
la fusion dans Releves-pv.csv, elle, reste dans data_loaders.py.

Detache de data_loaders.py le 28/09/2026 (1 900 lignes, six metiers) ;
data_loaders reimporte tout ce qui est ici : les appels existants
(dl.parse_enedis_fichier...) ne changent pas.
"""

from __future__ import annotations

import re
import statistics
from datetime import date as _date

from accords import accorde, pluriel

# ----------------------------------------------------------------------
# Import d'un export quotidien Enedis (consommation ou injection)
# ----------------------------------------------------------------------

_RE_DATE_ISO = re.compile(r"^(\d{4})-(\d{2})-(\d{2})")
_RE_DATE_FR = re.compile(r"^(\d{1,2})/(\d{1,2})/(\d{4})$")
# Une heure accolee a la date (« 2026-09-01T00:30 », « ... 00:30:00 »).
_RE_HEURE = re.compile(r"\d{1,2}:\d{2}")
_RE_NOMBRE = re.compile(r"^-?\d+(\.\d+)?$")


def _parse_date_any(s: str) -> str | None:
    """Date ISO (2026-06-01) ou FR (01/06/2026) -> 'DD/MM/YYYY', sinon None."""
    m = _RE_DATE_ISO.match(s)
    if m:
        y, mo, d = (int(x) for x in m.groups())
    else:
        m = _RE_DATE_FR.match(s)
        if not m:
            return None
        d, mo, y = (int(x) for x in m.groups())
    try:
        _date(y, mo, d)
    except ValueError:
        return None
    return f"{d:02d}/{mo:02d}/{y:04d}"


def _parse_nombre(s: str) -> float | None:
    s = s.strip().replace(" ", "").replace(" ", "").replace(",", ".")
    if not s or not _RE_NOMBRE.match(s):
        return None
    return float(s)


def _fmt_kwh_fr(v: float, decimales: int = 3) -> str:
    """12.571 -> '12,571' ; 9.0 -> '9' (format FR sans zeros inutiles)."""
    s = f"{v:.{decimales}f}".rstrip("0").rstrip(".")
    return s.replace(".", ",")


def _unite_annoncee(texte_bas: str) -> bool | None:
    """L'unite que l'en-tete annonce : True (kWh), False (Wh), None (rien).

    « (kWh) », « (en kWh) » ou « en kWh » ; la regle etait recopiee dans
    quatre lecteurs, chacun avec sa variante."""
    if "(kwh" in texte_bas or "en kwh" in texte_bas:
        return True
    if "(wh" in texte_bas or "en wh" in texte_bas:
        return False
    return None


def _noter_unite_devinee(notes: list | None, en_kwh: bool) -> None:
    """L'en-tete ne dit pas l'unite : elle est deduite de la grandeur des
    valeurs. On le dit dans le recapitulatif -- une erreur multiplierait
    tout par mille."""
    if notes is None:
        return
    if en_kwh:
        notes.append("Unité absente de l'en-tête : valeurs lues en kWh, "
                     "d'après leur grandeur. Vérifiez-les ci-dessus.")
    else:
        notes.append("Unité absente de l'en-tête : valeurs lues en Wh "
                     "(d'après leur grandeur), converties en kWh.")


def _noter_lignes_ecartees(notes: list | None, jours: list[str]) -> None:
    """Des lignes commencent par une date mais ne portent aucune valeur
    lisible : elles sont ignorees, et on le dit au lieu de se taire."""
    if notes is None or not jours:
        return
    n = len(jours)
    notes.append(f"{pluriel(n, 'ligne')} {accorde(n, 'datée')} sans valeur "
                 f"lisible, {accorde(n, 'ignorée')} (par exemple le "
                 f"{jours[0]}).")


def parse_enedis_csv(path: str) -> tuple[str, dict[str, str]]:
    """
    Lit un export quotidien Enedis (mon-compte-particulier.enedis.fr).

    Le format varie selon les versions du site : lignes de metadonnees,
    dates ISO ou FR, valeurs en Wh ou kWh. Tout est detecte automatiquement :
    - lignes retenues : celles qui commencent par une date suivie d'un nombre ;
    - type de donnees (via le texte du fichier) :
        'injection' / 'production'   -> colonne Inj_Jour
        'consommation' / 'soutirage' -> colonne Conso_réseau_Jour
    - unite : '(en kWh)' / '(en Wh)' dans les entetes, sinon heuristique
      (une valeur quotidienne mediane > 200 est forcement en Wh).

    Retourne (colonne_cible, {date 'DD/MM/YYYY': valeur kWh en chaine FR}).
    Leve ValueError avec un message clair si le fichier n'est pas reconnu.
    """
    return parse_enedis_csv_texte(_lire_texte(path))


_MOTS_INJECTION = ("injection", "production")
_MOTS_CONSO = ("consommation", "soutirage")


def _entete_du_fichier(texte: str) -> str:
    """Les lignes qui precedent la premiere ligne de donnees 'date;valeur'.

    C'est la que Enedis annonce le type ('Injection quotidienne'). Chercher
    dans tout le fichier serait trompeur : le nom de l'export contient
    souvent les DEUX mots ('Consommation-Production'), et les libelles
    parasites d'une autre section suffiraient a faire basculer le verdict.
    """
    entete = []
    for ligne in texte.splitlines():
        cellules = [c.strip().strip('"') for c in re.split(r"[;\t]", ligne)]
        if len(cellules) >= 2 and _parse_date_any(cellules[0]) is not None:
            if any(_parse_nombre(c) is not None for c in cellules[1:]):
                break  # premiere ligne de donnees : l'entete s'arrete ici
        entete.append(ligne)
    return "\n".join(entete)


def _type_de_donnees(texte: str) -> str:
    """Decide si l'export porte sur l'injection ou sur la consommation.

    Regle : on lit l'en-tete en priorite. Si les deux familles de mots y
    figurent, on REFUSE de deviner plutot que de risquer d'ecrire une conso
    dans la colonne Injection (l'ancienne version faisait gagner
    'production' sans autre forme de proces).
    """
    for zone, ou in ((_entete_du_fichier(texte), "l'en-tete"),
                     (texte, "le fichier")):
        bas = zone.lower()
        inj = [m for m in _MOTS_INJECTION if m in bas]
        conso = [m for m in _MOTS_CONSO if m in bas]
        if inj and conso:
            raise ValueError(
                f"Type de données ambigu : {ou} mentionne à la fois "
                f"{' / '.join(inj)} et {' / '.join(conso)}. Impossible de "
                "savoir s'il s'agit de l'injection ou de la consommation. "
                "Utilisez le classeur Excel Enedis, qui indique clairement "
                "chaque grandeur, ou un export de votre fournisseur qui le "
                "fait aussi (par exemple le CSV 'suivi de consommation' "
                "d'Octopus)."
            )
        if inj:
            return "Inj_Jour"
        if conso:
            return "Conso_réseau_Jour"
    raise ValueError(
        "Type de données introuvable : l'en-tête du fichier ne mentionne ni "
        "'injection'/'production' ni 'consommation', et l'application refuse "
        "de deviner dans quelle colonne ranger ces valeurs. Tout export "
        "quotidien 'date + valeur' convient -- Enedis, un fournisseur ou un "
        "onduleur -- pourvu que son en-tête dise de quelle grandeur il "
        "s'agit."
    )


def parse_enedis_csv_texte(texte: str, notes: list | None = None
                           ) -> tuple[str, dict[str, str]]:
    """Comme parse_enedis_csv, mais a partir du contenu deja lu.

    `notes` recoit ce que l'utilisateur doit savoir : unite devinee,
    lignes datees sans valeur lisible."""
    bas = texte.lower()
    colonne = _type_de_donnees(texte)
    # « Production » sans « injection » : chez Enedis c'est l'injection, mais
    # c'est aussi le mot d'un export d'onduleur passe par le mauvais bouton,
    # range alors en Injection sans un mot (audit du 08/10/2026). On le dit
    # dans le recapitulatif, qui s'affiche avant d'appliquer.
    if (notes is not None and colonne == "Inj_Jour"
            and "injection" not in _entete_du_fichier(texte).lower()):
        notes.append(
            "Ce fichier parle de « production » : il est rangé en Injection "
            "(ce que votre compteur envoie au réseau). Si c'est l'export de "
            "votre onduleur, cliquez sur Annuler et utilisez plutôt "
            "« Importer la production… ».")

    # Unite explicite dans les entetes ; None -> heuristique plus bas.
    unite_kwh = _unite_annoncee(bas)

    brutes: list[tuple[str, float]] = []
    dates_brutes: list[str] = []      # la cellule de date telle quelle
    ecartees: list[str] = []
    for ligne in texte.splitlines():
        # Point-virgule, tabulation ou virgule, comme le promet la Notice.
        cellules, virgule = _decoupe_ligne(ligne)
        if len(cellules) < 2:
            continue
        # Date seule, ou suivie d'une heure (« 01/06/2026 00:00 » : refuse
        # avant l'audit du 08/10/2026, alors que l'ISO avec heure passait).
        jour = _parse_jour_any(cellules[0])
        if jour is None:
            continue  # ligne de metadonnees / entete
        valeur = next(
            (v for v in (_nombre_cellule(c, virgule) for c in cellules[1:])
             if v is not None),
            None,
        )
        if valeur is None:
            # Une ligne 'date debut;date fin' des metadonnees n'est pas une
            # journee : on ne la compte pas parmi les lignes ecartees.
            if not any(_parse_date_any(c) for c in cellules[1:]):
                ecartees.append(jour)
            continue
        brutes.append((jour, valeur))
        dates_brutes.append(cellules[0])

    if not brutes:
        raise ValueError(
            "Aucune ligne 'date;valeur' reconnue. Attendu : un export "
            "quotidien Enedis (une date et un nombre par ligne, séparés "
            "par des points-virgules, des virgules ou des tabulations)."
        )

    # Un meme jour sur plusieurs lignes horodatees : c'est une courbe heure
    # par heure, pas un releve quotidien. La derniere valeur ecrasait les
    # autres en silence ; les additionner serait faux aussi, une courbe de
    # charge Enedis donnant souvent une puissance (W) et non une energie.
    # On refuse en disant quoi exporter (audit du 04/10/2026).
    horodates = [jour for (jour, _v), ligne_date in zip(brutes, dates_brutes,
                                                         strict=True)
                 if _RE_HEURE.search(ligne_date)]
    if len(horodates) != len(set(horodates)):
        raise ValueError(
            "Ce fichier donne plusieurs valeurs par jour (heure par heure ou "
            "à la demi-heure) : c'est une courbe de charge, pas un relevé "
            "quotidien. Téléchargez plutôt la consommation (ou l'injection) "
            "quotidienne, en kWh, une valeur par jour."
        )

    if unite_kwh is None:
        unite_kwh = statistics.median(v for _, v in brutes) < 200
        _noter_unite_devinee(notes, unite_kwh)
    _noter_lignes_ecartees(notes, ecartees)

    # Doublons internes au fichier importe : la derniere valeur gagne.
    valeurs: dict[str, str] = {}
    for jour, v in brutes:
        valeurs[jour] = _fmt_kwh_fr(v if unite_kwh else v / 1000.0)
    return colonne, valeurs


# ----------------------------------------------------------------------
# Import du "suivi de consommation" Octopus (detail heures creuses / pleines)
# ----------------------------------------------------------------------

# En-tetes du fichier "suivi_conso_<PRM>_<adresse>_<date>.csv", telecharge
# depuis l'espace client OCTOPUS -- et non chez Enedis, malgre le numero de
# PRM dans son nom : c'est le fournisseur qui l'edite, lui seul connaissant
# les prix (le fichier porte des colonnes en euros, qu'on ignore ici).
# C'etait longtemps le seul releve qui donnait le detail HC/HP jour par jour
# (le classeur .xlsx d'Enedis ne contient que le total) ; d'autres fournisseurs
# nomment ces colonnes autrement, et un utilisateur qui recopie ses chiffres a
# la main ecrit simplement "HC" et "HP". On accepte donc plusieurs libelles
# par colonne.
#
# La comparaison reste une EGALITE, apres normalisation : jamais un "contient",
# sinon "consommation hc (euros)" passerait pour des kWh.


def _normalise_entete(nom: str) -> str:
    """'  CONSO_HC ' -> 'conso hc' (minuscules, soulignements = espaces)."""
    return " ".join(nom.replace("_", " ").lower().split())


_ENTETES_CONSO = (
    "consommation (kwh)", "consommation", "conso (kwh)", "conso",
    "consommation totale (kwh)", "consommation totale",
    "consommation reseau (kwh)", "consommation réseau (kwh)",
    "conso reseau jour", "conso réseau jour", "conso reseau", "conso réseau",
)
_ENTETES_HC = (
    "consommation hc (kwh)", "consommation hc", "conso hc (kwh)", "conso hc",
    "hc (kwh)", "hc", "heures creuses (en kwh)",
    "heures creuses (kwh)", "heures creuses",
    "consommation heures creuses (kwh)", "consommation heures creuses",
)
_ENTETES_HP = (
    "consommation hp (kwh)", "consommation hp", "conso hp (kwh)", "conso hp",
    "hp (kwh)", "hp", "heures pleines (en kwh)",
    "heures pleines (kwh)", "heures pleines",
    "consommation heures pleines (kwh)", "consommation heures pleines",
)

_COLONNES_SUIVI_CONSO = {
    **{e: "Conso_réseau_Jour" for e in _ENTETES_CONSO},
    **{e: "Conso_HC" for e in _ENTETES_HC},
    **{e: "Conso_HP" for e in _ENTETES_HP},
}


def parse_octopus_suivi_conso(texte: str, notes: list | None = None
                              ) -> dict[str, dict[str, str]] | None:
    """
    Lit le "suivi de consommation" d'Octopus : une ligne par jour, avec des
    colonnes nommees dont le detail heures creuses / heures pleines.

    Les libelles acceptes pour chaque colonne sont listes dans _ENTETES_HC,
    _ENTETES_HP et _ENTETES_CONSO : le fichier d'Octopus, mais aussi des
    en-tetes simples ("HC", "HP", "Heures creuses"...) pour qui recopie ses
    chiffres a la main ou vient d'un autre fournisseur.

    Retourne {colonne_cible: {date 'DD/MM/YYYY': valeur kWh en chaine FR}},
    ou None si aucune colonne HC/HP n'est reconnue (le fichier releve alors
    du parseur generique parse_enedis_csv). Leve ValueError si une seule des
    deux colonnes est presente.
    """
    entetes: dict[int, str] | None = None
    resultat: dict[str, dict[str, str]] = {}
    ecartees: list[str] = []

    for ligne in texte.splitlines():
        # Meme decoupage que les autres lecteurs : sans lui, un fichier separe
        # par des tabulations n'etait pas reconnu, et le lecteur generique
        # rangeait alors les heures creuses en consommation totale.
        cellules, virgule = _decoupe_ligne(ligne)
        if len(cellules) < 2:
            continue

        if entetes is None:
            trouves = {
                i: _COLONNES_SUIVI_CONSO[_normalise_entete(nom)]
                for i, nom in enumerate(cellules)
                if _normalise_entete(nom) in _COLONNES_SUIVI_CONSO
            }
            cibles = set(trouves.values())
            # Sans le detail HC/HP, ce n'est pas ce format : on laisse la
            # main au parseur generique.
            if {"Conso_HC", "Conso_HP"} <= cibles:
                entetes = trouves
            elif cibles & {"Conso_HC", "Conso_HP"}:
                # Une seule des deux : le dire, plutot que de laisser le
                # lecteur generique ranger ces kWh en consommation reseau.
                if "Conso_HC" in cibles:
                    presente, manquante = "creuses", "pleines"
                    colonne_manquante = "HP"
                else:
                    presente, manquante = "pleines", "creuses"
                    colonne_manquante = "HC"
                raise ValueError(
                    f"Colonne {colonne_manquante} introuvable : ce fichier "
                    f"donne les heures {presente}, mais pas les heures "
                    f"{manquante}. Pour importer le détail, il faut les deux "
                    "colonnes, nommées par exemple 'Consommation HC (kWh)' "
                    "et 'Consommation HP (kWh)', ou simplement 'HC' et 'HP'."
                )
            continue

        jour = _parse_date_any(cellules[0])
        if jour is None:
            continue  # ligne de metadonnees ou de total
        lue = False
        for i, cible in entetes.items():
            if i >= len(cellules):
                continue
            valeur = _nombre_cellule(cellules[i], virgule)
            if valeur is not None:
                resultat.setdefault(cible, {})[jour] = _fmt_kwh_fr(valeur)
                lue = True
        if not lue:
            ecartees.append(jour)

    if resultat:
        _noter_lignes_ecartees(notes, ecartees)
    return resultat or None


# ----------------------------------------------------------------------
# Index de compteur -> consommations journalieres
# ----------------------------------------------------------------------

def _en_date(jour: str) -> _date:
    """'05/03/2026' -> date(2026, 3, 5)."""
    d, m, y = (int(x) for x in jour.split("/"))
    return _date(y, m, d)


def consommation_depuis_index(index: dict[str, float]) -> dict[str, str]:
    """
    Convertit des index de compteur (des cumuls) en consommations par jour.

    `index` : {date 'DD/MM/YYYY': index en kWh}. Un index ne fait que monter :
    la consommation d'une journee est la difference entre deux releves qui se
    suivent. L'index releve un matin cloturant la journee precedente, la
    difference entre le 02/01 et le 01/01 est rangee au 01/01.

    Deux cas sont ecartes plutot que devines :
    - un releve manquant (les deux dates ne se suivent pas d'un jour) : la
      difference couvrirait plusieurs journees, et rien ne dit comment la
      repartir. On laisse le trou ;
    - un index qui recule (compteur remplace, retour a zero) : la difference
      serait negative.

    Retourne {date: valeur kWh en chaine FR}, prete pour merge_import.
    """
    resultat: dict[str, str] = {}
    jours = sorted(index, key=_en_date)
    # strict=False : la deuxieme liste a un element de moins, c'est voulu.
    for precedent, suivant in zip(jours, jours[1:], strict=False):
        if (_en_date(suivant) - _en_date(precedent)).days != 1:
            continue
        ecart = index[suivant] - index[precedent]
        if ecart < 0:
            continue
        resultat[precedent] = _fmt_kwh_fr(ecart)
    return resultat


# ----------------------------------------------------------------------
# Export d'index quotidiens d'Enedis (.xlsx)
# ----------------------------------------------------------------------

# Telecharge dans l'espace client Enedis ("Export_<PRM>_Index_<periode>.xlsx").
# C'est le seul relevé du GESTIONNAIRE DE RESEAU qui donne les heures creuses
# et pleines : il est donc disponible a tout foyer francais, quel que soit son
# fournisseur. Mais il ne donne pas des consommations : ce sont les index du
# compteur, des cumuls -- d'ou consommation_depuis_index ci-dessus.
#
# Forme du fichier (relevee sur un vrai export, 18/09/2026) : quelques lignes
# de metadonnees, puis une ligne d'en-tetes "Date du tele-releve | Index
# totalisateur (en kWh) | Heures Pleines (en kWh) | Heures Creuses (en kWh) |
# Non parametre... | Heures Pleines Saison Basse (en kWh) | ...". Les colonnes
# HP/HC utiles sont celles du CALENDRIER FOURNISSEUR (l'offre de l'abonne) ;
# celles du CALENDRIER DISTRIBUTEUR, en saison basse / haute, decoupent la
# meme energie autrement : les confondre doublerait la consommation. D'ou une
# comparaison par EGALITE des libelles, jamais un "contient".
_ENTETES_INDEX_TOTAL = (
    "index totalisateur (en kwh)", "index totalisateur (kwh)",
    "index totalisateur",
)


def _charger_openpyxl():
    """Importe openpyxl, avec un message lisible s'il manque."""
    try:
        import openpyxl
    except ImportError as exc:
        raise ValueError(
            "La lecture d'un fichier Excel demande le module 'openpyxl' "
            "(installez-le avec : py -m pip install openpyxl)."
        ) from exc
    return openpyxl


def _index_d_une_feuille(ws, production: bool = False
                         ) -> tuple[dict[str, dict[str, str]] | None,
                                    str | None, str | None]:
    """
    Lit une feuille d'export d'index Enedis.

    `production` : cette feuille porte les index de production, c'est-a-dire
    l'energie INJECTEE au reseau (le compteur ne mesure que le surplus qui
    sort). Son en-tete est le meme que celui de la consommation : seul le nom
    de la feuille les distingue, d'ou ce drapeau -- sans lui, des kWh injectes
    seraient ranges en soutirage.

    Retourne ({colonne_cible: {date: kWh}} converti en consommations, premier
    jour lu, dernier jour lu), ou (None, None, None) si cette feuille n'est
    pas un export d'index.
    """
    total = "Inj_Jour" if production else "Conso_réseau_Jour"
    colonnes: dict[int, str] = {}
    col_date = -1
    index: dict[str, dict[str, float]] = {}

    for row in ws.iter_rows(values_only=True):
        cellules = list(row)

        if not colonnes:
            # Tant que l'en-tete n'est pas trouve, tout est metadonnees.
            trouves: dict[int, str] = {}
            date_ici = -1
            for i, valeur in enumerate(cellules):
                if not isinstance(valeur, str):
                    continue
                nom = _normalise_entete(valeur)
                if nom in _ENTETES_INDEX_TOTAL:
                    trouves[i] = total
                elif nom in _ENTETES_HP and not production:
                    trouves[i] = "Conso_HP"
                elif nom in _ENTETES_HC and not production:
                    trouves[i] = "Conso_HC"
                elif date_ici < 0 and nom.startswith("date"):
                    date_ici = i
            # L'index totalisateur est la signature du format : sans lui, ce
            # classeur est un autre export (conso/production), pas celui-ci.
            if date_ici >= 0 and total in trouves.values():
                colonnes, col_date = trouves, date_ici
            continue

        if col_date >= len(cellules):
            continue
        jour = _cellule_date(cellules[col_date])
        if jour is None:
            continue  # ligne vide, ou pied de tableau
        for i, cible in colonnes.items():
            if i >= len(cellules):
                continue
            valeur = _cellule_nombre(cellules[i])
            if valeur is not None:  # 'NA' et '-' : donnee indisponible
                index.setdefault(cible, {})[jour] = valeur

    if not colonnes:
        return None, None, None
    # La consommation du jour J est l'index du lendemain moins celui de J :
    # le premier jour du fichier a donc sa valeur, mais le dernier attend
    # l'index du lendemain. L'appelant le dit a l'utilisateur, qui sinon
    # cherche en vain ce jour dans ses releves.
    tous = [j for valeurs in index.values() for j in valeurs]
    premier = min(tous, key=_en_date) if tous else None
    dernier = max(tous, key=_en_date) if tous else None
    return ({cible: consommation_depuis_index(valeurs)
             for cible, valeurs in index.items()}, premier, dernier)


def parse_enedis_index_xlsx(path: str, notes: list | None = None
                            ) -> dict[str, dict[str, str]] | None:
    """
    Lit un export d'index quotidiens d'Enedis et le convertit en
    consommations journalieres. Chez un producteur, le classeur porte aussi
    une feuille d'index de production : elle donne l'injection.

    Retourne None si le classeur n'est pas de ce type : l'appelant passe
    alors au lecteur du classeur conso/production habituel.

    `notes` : liste que l'appelant fournit s'il veut afficher les remarques
    de lecture -- ici, le dernier jour, qui attend l'index du lendemain.
    """
    openpyxl = _charger_openpyxl()
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    resultat: dict[str, dict[str, str]] = {}
    reconnu = False
    try:
        # Un producteur a deux feuilles : la consommation (avec le detail
        # HP/HC) et la production, c'est-a-dire l'injection. On prend les
        # deux, chacune dans sa colonne.
        derniers = []
        for nom in wb.sheetnames:
            trouve, _premier, dernier = _index_d_une_feuille(
                wb[nom], "prod" in nom.lower())
            if trouve is None:
                continue  # feuille de garde, ou autre chose
            reconnu = True
            resultat.update(trouve)
            if dernier:
                derniers.append(dernier)
    finally:
        wb.close()
    if not reconnu:
        return None
    # Le premier jour de l'export, lui, a bien sa valeur (index du lendemain
    # moins le sien) : aucune remarque a faire a son sujet.
    if notes is not None and derniers:
        dernier = max(derniers, key=_en_date)
        notes.append(
            f"Le {dernier} n'apparaît pas encore : sa consommation se "
            "calcule avec le relevé du lendemain, qu'Enedis publie un "
            "jour plus tard. Il se complétera au prochain import, si "
            f"l'export part du {dernier} ou d'avant.")
    return resultat


# Feuilles reconnues dans le classeur Excel Enedis. Pour un producteur,
# la "production" mesuree par Enedis est l'energie injectee au reseau.
_FEUILLES_ENEDIS = (
    ("production", "Inj_Jour"),
    ("injection", "Inj_Jour"),
    ("consommation", "Conso_réseau_Jour"),
    ("soutirage", "Conso_réseau_Jour"),
)


def _cellule_date(c) -> str | None:
    """Cellule Excel -> 'DD/MM/YYYY' si c'est une date, sinon None."""
    if isinstance(c, _date):  # datetime.datetime herite de datetime.date
        return f"{c.day:02d}/{c.month:02d}/{c.year:04d}"
    if isinstance(c, str):
        return _parse_date_any(c.strip())
    return None


def _cellule_nombre(c) -> float | None:
    if isinstance(c, (int, float)) and not isinstance(c, bool):
        return float(c)
    if isinstance(c, str):
        return _parse_nombre(c)
    return None


def parse_enedis_xlsx(path: str, notes: list | None = None
                      ) -> dict[str, dict[str, str]]:
    """
    Lit un export Excel Enedis ('..._Export_energie_Consommation-Production_...').

    Le classeur contient une feuille par grandeur ('Export Consommation
    Quotidienne', 'Export Production Quotidienne') : des lignes de
    metadonnees puis un tableau 'Date | Valeur (en kWh)'. Les jours 'NA'
    (donnees indisponibles) sont ignores.

    Retourne {colonne_cible: {date 'DD/MM/YYYY': valeur kWh en chaine FR}}.
    """
    try:
        import openpyxl
    except ImportError as exc:
        raise ValueError(
            "La lecture d'un fichier Excel demande le module 'openpyxl' "
            "(installez-le avec : py -m pip install openpyxl)."
        ) from exc

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    resultat: dict[str, dict[str, str]] = {}
    try:
        for nom in wb.sheetnames:
            colonne = next(
                (col for cle, col in _FEUILLES_ENEDIS if cle in nom.lower()),
                None)
            if colonne is None:
                continue  # ex. "Page d'accueil"
            brutes: list[tuple[str, float]] = []
            textes: list[str] = []
            ecartees: list[str] = []
            for row in wb[nom].iter_rows(values_only=True):
                cellules = [c for c in row if c is not None]
                textes.extend(c for c in cellules if isinstance(c, str))
                if len(cellules) < 2:
                    continue
                jour = _cellule_date(cellules[0])
                if jour is None:
                    continue  # metadonnees ('Date de debut : ', ...)
                valeur = next(
                    (v for v in (_cellule_nombre(c) for c in cellules[1:])
                     if v is not None), None)
                if valeur is not None:
                    brutes.append((jour, valeur))
                else:
                    ecartees.append(jour)   # 'NA' : donnee indisponible
            if not brutes:
                continue
            bas = " ".join(textes).lower()
            en_kwh = _unite_annoncee(bas)
            if en_kwh is None:
                en_kwh = statistics.median(v for _, v in brutes) < 200
                _noter_unite_devinee(notes, en_kwh)
            _noter_lignes_ecartees(notes, ecartees)
            valeurs: dict[str, str] = {}
            for jour, v in brutes:  # doublon interne : la derniere gagne
                valeurs[jour] = _fmt_kwh_fr(v if en_kwh else v / 1000.0)
            resultat[colonne] = valeurs
    finally:
        wb.close()

    if not resultat:
        raise ValueError(
            "Aucune feuille 'Consommation' ou 'Production' avec des lignes "
            "'date + valeur' dans ce classeur. Attendu : le classeur Excel "
            "d'un gestionnaire de réseau (Enedis) ou un export équivalent."
        )
    return resultat


# Export mensuel du portail d'un onduleur ("rapport de centrale" chez Huawei
# FusionSolar, formes voisines ailleurs) : une ligne par jour, une vingtaine
# de colonnes en kWh.
#
# Il apporte ce qu'aucun releve Enedis ne peut donner AVANT la mise en service
# du compteur : ce qui a ete produit, injecte et soutire pendant ces journees.
# Sans lui, l'application doit estimer l'autoconsommation de cette periode --
# des semaines, parfois une saison entiere chez qui attend son raccordement.
#
# Trois colonnes seulement sont importees. La "Consommation" et
# l'"Autoconsommation" du rapport sont redondantes : l'application les
# recalcule (autoconso = production - injection). Elles servent de controle.
#
# Piege principal : "Production totale" est un CUMUL qui ne repart jamais a
# zero (88,16 le 31/01 puis 129,46 le 01/02). La prendre pour la production du
# jour multiplierait les chiffres par dix en fin de mois. C'est "Production PV"
# qui donne le jour, et elle seule est lue.
_COLONNES_RAPPORT_ONDULEUR = {
    "production pv (kwh)": "Prod_Jour",
    "production pv": "Prod_Jour",
    "exportation (kwh)": "Inj_Jour",
    "exportation": "Inj_Jour",
    "injection (kwh)": "Inj_Jour",
    "importation (kwh)": "Conso_réseau_Jour",
    "importation": "Conso_réseau_Jour",
}
_ENTETES_AUTOCONSO_RAPPORT = ("autoconsommation (kwh)", "autoconsommation")
_ENTETES_DATE_RAPPORT = (
    "période statistique", "periode statistique", "date statistique",
    "date", "jour",
)


def _rapport_d_une_feuille(ws, notes: list | None
                           ) -> dict[str, dict[str, str]] | None:
    """Lit une feuille de rapport d'onduleur, ou None si ce n'en est pas une."""
    cibles: dict[int, str] = {}
    col_date: int | None = None
    col_autoconso: int | None = None
    entetes_bas = ""
    lignes_donnees: list[tuple] = []

    for row in ws.iter_rows(values_only=True):
        if not cibles:
            # Tant qu'on n'a pas reconnu l'en-tete, chaque ligne est un
            # candidat : le titre du rapport occupe la premiere.
            vus: dict[int, str] = {}
            for i, cellule in enumerate(row):
                if not isinstance(cellule, str):
                    continue
                nom = _normalise_entete(cellule)
                if nom in _COLONNES_RAPPORT_ONDULEUR:
                    vus[i] = _COLONNES_RAPPORT_ONDULEUR[nom]
                elif nom in _ENTETES_AUTOCONSO_RAPPORT:
                    col_autoconso = i
                elif nom in _ENTETES_DATE_RAPPORT and col_date is None:
                    col_date = i
            # Deux grandeurs reconnues au moins : une seule pourrait etre
            # une coincidence d'en-tete dans un tout autre classeur.
            if len(set(vus.values())) >= 2:
                cibles = vus
                entetes_bas = " ".join(
                    c.lower() for c in row if isinstance(c, str))
                if col_date is None:
                    col_date = 0
            else:
                col_autoconso = None
                col_date = None
            continue

        jour = _cellule_date(row[col_date]) if col_date < len(row) else None
        if jour is None:
            continue  # ligne vide, ou total de bas de tableau
        lignes_donnees.append((jour, row))

    if not cibles or not lignes_donnees:
        return None

    # Unite : l'en-tete l'annonce presque toujours ("(kWh)"). Sinon, meme
    # repere que pour les autres imports -- une production journaliere
    # domestique depasse toujours 200 en Wh et reste en dessous en kWh.
    en_kwh = _unite_annoncee(entetes_bas)
    if en_kwh is None:
        prods = [v for v in (_valeur_ligne(r, i) for _j, r in lignes_donnees
                             for i, c in cibles.items() if c == "Prod_Jour")
                 if v is not None]
        en_kwh = not prods or statistics.median(prods) < 200
        if prods:
            _noter_unite_devinee(notes, en_kwh)

    resultat: dict[str, dict[str, str]] = {}
    ecarts = comparables = 0
    for jour, row in lignes_donnees:
        du_jour: dict[str, float] = {}
        for i, colonne in cibles.items():
            valeur = _valeur_ligne(row, i)
            if valeur is None:
                continue
            du_jour[colonne] = valeur if en_kwh else valeur / 1000.0
            resultat.setdefault(colonne, {})[jour] = _fmt_kwh_fr(du_jour[colonne])
        # Controle : l'autoconsommation annoncee doit valoir
        # production - injection. Si elle ne tombe pas juste, les colonnes
        # lues ne sont pas celles qu'on croit.
        auto = _valeur_ligne(row, col_autoconso) if col_autoconso is not None else None
        if (auto is not None and "Prod_Jour" in du_jour and "Inj_Jour" in du_jour):
            comparables += 1
            attendu = du_jour["Prod_Jour"] - du_jour["Inj_Jour"]
            if abs((auto if en_kwh else auto / 1000.0) - attendu) > 0.05:
                ecarts += 1

    if not resultat:
        return None
    if notes is not None and comparables and ecarts > comparables * 0.10:
        notes.append(
            f"Attention : sur {pluriel(comparables, 'jour')}, {ecarts} "
            f"{accorde(ecarts, 'a', 'ont')} une "
            "autoconsommation qui ne vaut pas « production − exportation ». "
            "Les colonnes lues ne sont peut-être pas les bonnes : vérifiez "
            "quelques journées avant d'appliquer.")
    return resultat


def _valeur_ligne(row, i: int | None) -> float | None:
    """Valeur numerique de la colonne i d'une ligne, ou None."""
    if i is None or i >= len(row):
        return None
    return _cellule_nombre(row[i])


def parse_rapport_onduleur_xlsx(path: str, notes: list | None = None
                                ) -> dict[str, dict[str, str]] | None:
    """
    Lit le rapport mensuel du portail d'un onduleur (Huawei FusionSolar et
    equivalents) : une ligne par jour, production / exportation / importation.

    Retourne {colonne_cible: {date 'DD/MM/YYYY': kWh en chaine FR}}, ou None
    si le classeur n'est pas de ce format -- l'appelant essaie alors un
    autre lecteur.
    """
    openpyxl = _charger_openpyxl()
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for nom in wb.sheetnames:
            trouve = _rapport_d_une_feuille(wb[nom], notes)
            if trouve is not None:
                return trouve
    finally:
        wb.close()
    return None


def parse_enedis_fichier(path: str, notes: list | None = None
                         ) -> dict[str, dict[str, str]]:
    """
    Point d'entree de l'import d'un releve du reseau, trois formats acceptes :
    - .xlsx : le classeur officiel d'Enedis (feuilles conso + production),
      ou son export d'index quotidiens (converti en consommations) ;
    - .csv "suivi de consommation" d'Octopus : conso + detail HC/HP ;
    - .csv simple : une seule grandeur (conso ou injection).

    Le nom garde "enedis" pour ne pas casser les appels existants : les deux
    fournisseurs passent par ici.

    Retourne {colonne_cible: {date 'DD/MM/YYYY': valeur kWh en chaine FR}}.
    """
    if path.lower().endswith((".xlsx", ".xlsm")):
        index = parse_enedis_index_xlsx(path, notes)
        if index is not None:
            return index
        return parse_enedis_xlsx(path, notes)
    texte = _lire_texte(path)
    suivi = parse_octopus_suivi_conso(texte, notes)
    if suivi is not None:
        return suivi
    colonne, valeurs = parse_enedis_csv_texte(texte, notes)
    return {colonne: valeurs}


# ----------------------------------------------------------------------
# Import d'un rapport Enphase (production)
# ----------------------------------------------------------------------

# Date suivie eventuellement d'une heure : "2026-07-14 06:15:00 +0200",
# "14/07/2026 06:15". On ne garde que le jour, les valeurs sont totalisees.
# Annee sur quatre ou deux chiffres dans un fichier importe (« 01/06/26 » :
# 2026 ; audit du 08/10/2026, la ligne etait ignoree sans un mot). La saisie
# a la main, elle, reste stricte (data_loaders.normalise_date_saisie).
_RE_DATE_HEURE_FR = re.compile(r"^(\d{1,2})/(\d{1,2})/(\d{4}|\d{2})(?!\d)")


def _parse_jour_any(s: str) -> str | None:
    """Comme _parse_date_any mais accepte une heure derriere la date, et une
    annee sur deux chiffres (fichiers importes seulement)."""
    jour = _parse_date_any(s)
    if jour is not None:
        return jour
    m = _RE_DATE_HEURE_FR.match(s.strip())
    if not m:
        return None
    d, mo, y = (int(x) for x in m.groups())
    if y < 100:
        y += 2000
    try:
        _date(y, mo, d)
    except ValueError:
        return None
    return f"{d:02d}/{mo:02d}/{y:04d}"


def _decoupe_ligne(ligne: str) -> tuple[list[str], bool]:
    """Decoupe une ligne CSV. Retourne (cellules, separateur_virgule)."""
    if ";" in ligne:
        sep, virgule = ";", False
    elif "\t" in ligne:
        sep, virgule = "\t", False
    elif "," in ligne:
        sep, virgule = ",", True
    else:
        return [], False
    return [c.strip().strip('"') for c in ligne.split(sep)], virgule


def _nombre_cellule(cell: str, virgule: bool) -> float | None:
    """Lit un nombre dans une cellule decoupee par _decoupe_ligne.

    Dans un fichier separe par des virgules, seul le point est decimal
    ("5.5") ; sinon la virgule francaise est acceptee ("5,5")."""
    if not virgule:
        return _parse_nombre(cell)
    brut = cell.replace(" ", "").replace(" ", "")
    return float(brut) if _RE_NOMBRE.match(brut) else None


def parse_enphase_texte(texte: str, notes: list | None = None
                        ) -> dict[str, str]:
    """
    Lit le contenu d'un rapport Enphase Enlighten (menu Systeme > Rapports,
    "Energie mensuelle" recu par email).

    Le rapport donne l'energie produite toutes les 15 minutes : les valeurs
    d'un meme jour sont **totalisees** pour obtenir la production du jour.
    Un rapport deja quotidien (une ligne par jour) fonctionne aussi.

    Tolerant sur la forme : lignes d'entete ignorees, separateur ';', ',' ou
    tabulation, dates ISO ou FR avec ou sans heure, valeurs en Wh ou kWh
    (unite lue dans l'entete, sinon devinee sur le total journalier).

    Retourne {date 'DD/MM/YYYY': production kWh en chaine FR}, arrondie au
    dixieme comme les valeurs affichees par Enlighten et le reste du CSV.
    """
    bas = texte.lower()
    unite_kwh = _unite_annoncee(bas)

    totaux: dict[str, float] = {}
    ecartees: list[str] = []
    for ligne in texte.splitlines():
        cellules, virgule = _decoupe_ligne(ligne)
        if len(cellules) < 2:
            continue
        jour = _parse_jour_any(cellules[0])
        if jour is None:
            continue  # entete ou ligne de metadonnees
        valeur = next(
            (v for v in (_nombre_cellule(c, virgule) for c in cellules[1:])
             if v is not None),
            None,
        )
        if valeur is None:
            if jour not in ecartees:
                ecartees.append(jour)
            continue
        totaux[jour] = totaux.get(jour, 0.0) + valeur

    if not totaux:
        raise ValueError(
            "Aucune ligne 'date;valeur' reconnue. Attendu : la production "
            "journalière de votre onduleur -- par exemple le rapport "
            "Enphase 'Energie mensuelle' (Enlighten > Menu > Systeme > "
            "Rapports), ou le rapport d'onduleur au format Excel."
        )

    if unite_kwh is None:
        # Un total journalier de production depasse toujours 200 en Wh et
        # reste bien en dessous en kWh (installation domestique).
        unite_kwh = statistics.median(totaux.values()) < 200
        _noter_unite_devinee(notes, unite_kwh)
    _noter_lignes_ecartees(notes, ecartees)

    return {
        jour: _fmt_kwh_fr(v if unite_kwh else v / 1000.0, decimales=1)
        for jour, v in totaux.items()
    }


def parse_enphase_fichier(path: str, notes: list | None = None
                          ) -> dict[str, dict[str, str]]:
    """
    Point d'entree de l'import de production : fichier CSV/texte du rapport
    Enphase, archive .zip telle qu'elle arrive dans l'email (tous ses
    fichiers .csv sont lus, voir _lire_archive), ou classeur Excel du portail
    d'un onduleur ("rapport de centrale" Huawei et equivalents).

    Retourne {'Prod_Jour': {date 'DD/MM/YYYY': valeur kWh en chaine FR}} ;
    le rapport d'onduleur y ajoute l'injection et la consommation reseau,
    qu'il est seul a connaitre avant la mise en service du compteur.
    """
    if path.lower().endswith((".xlsx", ".xlsm")):
        rapport = parse_rapport_onduleur_xlsx(path, notes)
        if rapport is None:
            raise ValueError(
                "Ce classeur Excel n'a pas la forme d'un rapport d'onduleur "
                "(une ligne par jour, des colonnes « Production PV », "
                "« Exportation », « Importation »).\n\nS'il vient de votre "
                "gestionnaire de réseau, importez-le avec le bouton "
                "« Importer achat au réseau / injection… ».")
        return rapport

    if path.lower().endswith(".zip"):
        return {"Prod_Jour": _lire_archive(path, notes)}
    with open(path, "rb") as fh:
        texte = _decode_texte(fh.read())
    return {"Prod_Jour": parse_enphase_texte(texte, notes)}


def _lire_archive(path: str, notes: list | None) -> dict[str, str]:
    """Lit TOUS les releves d'une archive .zip et les reunit par date.

    Seul le premier fichier etait lu : les mois suivants d'une archive
    disparaissaient sans un mot. Les .csv d'abord (un lisez-moi .txt ne
    nous interesse pas) ; les .txt seulement s'il n'y a aucun .csv. Deux
    fichiers qui donnent deux productions differentes pour le meme jour :
    on s'arrete et on le dit, plutot que de choisir au hasard."""
    import zipfile

    with zipfile.ZipFile(path) as zf:
        noms = sorted(n for n in zf.namelist() if n.lower().endswith(".csv"))
        if not noms:
            noms = sorted(n for n in zf.namelist()
                          if n.lower().endswith(".txt"))
        if not noms:
            raise ValueError(
                "Cette archive ne contient aucun fichier .csv. "
                "Attendu : un relevé de production, par exemple le "
                "rapport Enphase reçu par e-mail.")
        contenus = [(nom, _decode_texte(zf.read(nom))) for nom in noms]

    valeurs: dict[str, str] = {}
    origine: dict[str, str] = {}
    for nom, texte in contenus:
        for jour, v in parse_enphase_texte(texte, notes).items():
            if jour in valeurs and valeurs[jour] != v:
                raise ValueError(
                    f"Les fichiers « {origine[jour]} » et « {nom} » de "
                    f"l'archive donnent deux productions différentes pour "
                    f"le {jour} ({valeurs[jour]} et {v}\xa0kWh). Importez-les "
                    "un par un, en commençant par celui qui fait foi.")
            valeurs[jour] = v
            origine.setdefault(jour, nom)
    if notes is not None and len(noms) > 1:
        notes.append(f"Archive : {len(noms)} fichiers lus "
                     f"({', '.join(noms)}).")
    return valeurs


def _decode_texte(brut: bytes) -> str:
    """Decode un contenu texte en essayant les encodages usuels.

    Le « Texte Unicode » de Windows (UTF-16) se reconnait a sa marque de
    tete, ou, sans elle, aux octets nuls entre les lettres : decode en
    UTF-8, il passait sans erreur mais illisible (« type de données
    introuvable », audit du 08/10/2026)."""
    if brut.startswith((b"\xff\xfe", b"\xfe\xff")):
        return brut.decode("utf-16")
    debut = brut[:200]
    if debut and debut.count(b"\x00") > len(debut) // 4:
        return brut.decode("utf-16-be" if debut[0] == 0 else "utf-16-le")
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return brut.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("Encodage du fichier non reconnu.")


def _lire_texte(path: str) -> str:
    """Lit un fichier texte quel que soit son encodage."""
    with open(path, "rb") as fh:
        return _decode_texte(fh.read())
